"""Generación de pautas de respuestas para formas randomizadas.

Funciones:
- parse_answer_key_csv: Parsea la pauta original (CSV exportado de Google Sheets)
- generate_answer_key: Genera el mapeo {pregunta → respuesta} para una forma randomizada
- build_answer_key_xlsx: Escribe la pauta como archivo .xlsx
"""

import csv
import io

from openpyxl import Workbook

from src.models import Question


def parse_answer_key_csv(csv_content: str) -> dict[int, str]:
    """Parsea el contenido CSV de la pauta original.

    Formato esperado:
        Pregunta,Respuesta
        1,a
        2,b
        ...

    Args:
        csv_content: String con el contenido CSV completo.

    Returns:
        Dict mapeando número de pregunta a letra de respuesta correcta.
        Ejemplo: {1: "a", 2: "b", 3: "c"}
    """
    reader = csv.DictReader(io.StringIO(csv_content))
    answer_key: dict[int, str] = {}

    for row in reader:
        q_num_raw = row.get("Pregunta", "").strip()
        answer_raw = row.get("Respuesta", "").strip().lower()

        if not q_num_raw or not answer_raw:
            continue

        answer_key[int(q_num_raw)] = answer_raw

    return answer_key


def generate_answer_key(
    questions: list[Question],
    original_key: dict[int, str],
) -> dict[int, str]:
    """Genera la pauta de respuestas para una forma randomizada.

    Para cada pregunta:
    1. Buscar la respuesta correcta original usando original_number
    2. Encontrar qué alternativa tiene esa original_letter
    3. La letra actual de esa alternativa es la nueva respuesta correcta

    Args:
        questions: Lista de preguntas randomizadas (con original_number/original_letter).
        original_key: Pauta original {número_pregunta: letra_correcta}.

    Returns:
        Dict {nuevo_número: nueva_letra_correcta}.

    Raises:
        KeyError: Si original_number no existe en original_key.
        ValueError: Si la letra correcta original no se encuentra en las alternativas.
    """
    new_key: dict[int, str] = {}

    for q in questions:
        correct_original_letter = original_key[q.original_number]

        new_letter = None
        for alt in q.alternatives:
            if alt.original_letter == correct_original_letter:
                new_letter = alt.letter
                break

        if new_letter is None:
            raise ValueError(
                f"Q{q.number} (original Q{q.original_number}): "
                f"no se encontró alternativa con original_letter='{correct_original_letter}'. "
                f"Alternativas: {[(a.letter, a.original_letter) for a in q.alternatives]}"
            )

        new_key[q.number] = new_letter

    return new_key


def build_answer_key_xlsx(
    answer_key: dict[int, str],
    output_path: str,
) -> None:
    """Escribe la pauta de respuestas como archivo .xlsx.

    Formato:
        A1: "Pregunta", B1: "Respuesta"
        A2: 1, B2: "a"
        A3: 2, B3: "b"
        ...

    Las filas se ordenan por número de pregunta ascendente.

    Args:
        answer_key: Dict {número_pregunta: letra_correcta}.
        output_path: Ruta donde guardar el .xlsx.
    """
    wb = Workbook()
    ws = wb.active

    ws.cell(row=1, column=1, value="Pregunta")
    ws.cell(row=1, column=2, value="Respuesta")

    for row_idx, q_num in enumerate(sorted(answer_key.keys()), start=2):
        ws.cell(row=row_idx, column=1, value=q_num)
        ws.cell(row=row_idx, column=2, value=answer_key[q_num])

    wb.save(output_path)
