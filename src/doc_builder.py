"""Generador de archivos .docx por clonación del documento original.

En lugar de recrear el documento desde cero, este módulo:
1. Copia el archivo .docx original byte a byte
2. Abre la copia con python-docx
3. Reemplaza el contenido de cada celda de la tabla de preguntas
   con las preguntas randomizadas (preservando XML/formato/imágenes)
4. Actualiza el texto "Forma X" en el encabezado
5. Convierte imágenes flotantes (anchor) a inline para que fluyan con el texto
6. Escala imágenes que excedan el ancho de celda

Esto garantiza que TODO el formato, estilos, imágenes, logos, tablas
de datos e instrucciones se preservan exactamente como en el original.
"""

import re
import shutil
from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.models import Question

NSMAP = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

# Namespaces for image handling
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"

# Default max image width in EMU (3.2 inches) — conservative fit for ~3.5" cells
# Used as fallback when column width can't be read from the table grid.
# 1 inch = 914400 EMU
DEFAULT_MAX_IMAGE_WIDTH_EMU = int(3.2 * 914400)  # 2,926,080 EMU

# Default cell padding in EMU (~115 dxa per side × 2 sides = 230 dxa)
# Subtracted from column width to get usable image width.
# 1 dxa = 635 EMU
CELL_PADDING_EMU = int(230 * 635)  # 146,050 EMU


def _clear_cell(cell):
    """Elimina todos los párrafos de una celda excepto las propiedades.

    Las celdas de Word DEBEN tener al menos un párrafo (según el estándar),
    pero nosotros las vaciamos y luego insertamos los nuevos párrafos.

    Args:
        cell: Objeto Cell de python-docx.
    """
    tc = cell._tc
    # Remover todos los <w:p> existentes
    for p in tc.findall("w:p", NSMAP):
        tc.remove(p)


def _add_paragraphs_to_cell(cell, paragraphs: list):
    """Agrega una lista de elementos <w:p> XML a una celda.

    Args:
        cell: Objeto Cell de python-docx.
        paragraphs: Lista de elementos lxml <w:p> a insertar.
    """
    tc = cell._tc
    for p_elem in paragraphs:
        tc.append(deepcopy(p_elem))


def _paragraph_has_image(p_element) -> bool:
    """Detecta si un párrafo contiene una imagen (drawing o AlternateContent).

    Args:
        p_element: Elemento XML <w:p>.

    Returns:
        True si el párrafo contiene un elemento de imagen.
    """
    if p_element.findall(f".//{{{NS_W}}}drawing"):
        return True
    if p_element.findall(f".//{{{NS_MC}}}AlternateContent"):
        return True
    return False


def _paragraph_has_text(p_element) -> bool:
    """Detecta si un párrafo contiene texto visible (al menos un <w:t> con contenido).

    Args:
        p_element: Elemento XML <w:p>.

    Returns:
        True si el párrafo tiene texto no vacío.
    """
    for t in p_element.findall(f".//{{{NS_W}}}t"):
        if t.text and t.text.strip():
            return True
    return False


def _split_text_and_image_runs(p_element) -> list:
    """Separa un párrafo mixto (texto + imagen en el mismo párrafo) en dos párrafos.

    Cuando un párrafo tiene tanto runs de texto como runs con <w:drawing>,
    los coloca en el mismo bloque y la imagen queda adherida al texto sin
    espacio. Esta función los divide en:
      1. Un párrafo con solo los runs de texto (preservando <w:pPr>).
      2. Un párrafo con solo los runs de imagen (preservando <w:pPr>).

    Si el párrafo solo tiene texto o solo tiene imagen, lo retorna tal cual
    (lista de un solo elemento sin copiar, para evitar trabajo innecesario).

    Args:
        p_element: Elemento XML <w:p> — puede ser mixto o puro.

    Returns:
        Lista con uno o dos elementos <w:p> lxml.
    """
    has_img = _paragraph_has_image(p_element)
    has_txt = _paragraph_has_text(p_element)

    # Párrafo puro (solo texto o solo imagen) → sin cambios
    if not (has_img and has_txt):
        return [p_element]

    # Obtener <w:pPr> para copiarlo en ambos párrafos resultantes
    pPr = p_element.find(f"{{{NS_W}}}pPr")

    # Clasificar runs: los que tienen <w:drawing> van al párrafo de imagen,
    # el resto (texto, rPr-only, etc.) van al párrafo de texto.
    text_runs = []
    image_runs = []
    for child in p_element:
        tag = child.tag
        if tag == f"{{{NS_W}}}pPr":
            continue  # Se copia aparte en ambos párrafos
        if tag == f"{{{NS_W}}}r":
            if child.find(f"{{{NS_W}}}drawing") is not None:
                image_runs.append(child)
            else:
                text_runs.append(child)
        else:
            # Otros elementos (bookmarkStart, proofErr, etc.) van al párrafo de texto
            text_runs.append(child)

    from lxml import etree

    def _make_paragraph_clone(runs: list) -> etree._Element:
        """Crea un nuevo <w:p> con los runs dados y una copia de <w:pPr>."""
        new_p = etree.Element(f"{{{NS_W}}}p")
        if pPr is not None:
            new_p.append(deepcopy(pPr))
        for run in runs:
            new_p.append(deepcopy(run))
        return new_p

    result = []
    if text_runs:
        result.append(_make_paragraph_clone(text_runs))
    if image_runs:
        result.append(_make_paragraph_clone(image_runs))

    # Fallback: si por alguna razón quedó vacío, retornar el original
    return result if result else [p_element]


def _convert_anchors_to_inline(p_element):
    """Convierte imágenes flotantes (<wp:anchor>) a inline (<wp:inline>).

    Las imágenes anchor tienen posición absoluta (relativa a columna/párrafo).
    Cuando las preguntas se reordenan, el texto cambia de longitud pero la
    posición absoluta de la imagen no, causando superposición con el texto.

    Las imágenes inline fluyen con el texto — se posicionan donde están en el
    flujo del documento, como un carácter más. Esto elimina el problema.

    La conversión:
    1. Cambia el tag de <wp:anchor> a <wp:inline>
    2. Remueve atributos exclusivos de anchor (posición, overlap, etc.)
    3. Remueve elementos hijos exclusivos de anchor (positionH/V, wrap*, simplePos)
    4. Mantiene extent, effectExtent, docPr y graphic (compartidos)

    Args:
        p_element: Elemento XML <w:p> que puede contener imágenes anchor.
    """
    for anchor in p_element.findall(f".//{{{NS_WP}}}anchor"):
        # 1. Cambiar el tag de anchor a inline
        anchor.tag = f"{{{NS_WP}}}inline"

        # 2. Remover atributos exclusivos de anchor
        anchor_only_attrs = (
            "allowOverlap",
            "behindDoc",
            "hidden",
            "layoutInCell",
            "locked",
            "relativeHeight",
            "simplePos",
        )
        for attr in anchor_only_attrs:
            if attr in anchor.attrib:
                del anchor.attrib[attr]

        # 3. Remover elementos hijos exclusivos de anchor
        anchor_only_children = (
            "simplePos",
            "positionH",
            "positionV",
            "wrapNone",
            "wrapSquare",
            "wrapTight",
            "wrapThrough",
            "wrapTopAndBottom",
        )
        for child_tag in anchor_only_children:
            for child in anchor.findall(f"{{{NS_WP}}}{child_tag}"):
                anchor.remove(child)


def _clamp_image_width(p_element, max_width_emu: int):
    """Escala imágenes que excedan el ancho máximo, preservando aspect ratio.

    Actualiza tanto <wp:extent> (tamaño del contenedor inline/anchor) como
    <a:ext> dentro de <pic:spPr>/<a:xfrm> (tamaño renderizado de la imagen).

    Args:
        p_element: Elemento XML <w:p> que puede contener imágenes.
        max_width_emu: Ancho máximo permitido en EMUs (English Metric Units).
    """
    for extent in p_element.findall(f".//{{{NS_WP}}}extent"):
        cx = int(extent.get("cx", "0"))
        cy = int(extent.get("cy", "0"))

        if cx <= max_width_emu or cx <= 0:
            continue

        scale = max_width_emu / cx
        new_cx = max_width_emu
        new_cy = int(cy * scale)

        extent.set("cx", str(new_cx))
        extent.set("cy", str(new_cy))

        # También actualizar <a:ext> dentro del mismo inline/anchor
        # (el tamaño renderizado de la forma de la imagen)
        container = extent.getparent()
        if container is not None:
            for a_ext in container.findall(f".//{{{NS_A}}}ext"):
                a_cx = int(a_ext.get("cx", "0"))
                a_cy = int(a_ext.get("cy", "0"))
                if a_cx > 0:
                    a_ext.set("cx", str(int(a_cx * scale)))
                    a_ext.set("cy", str(int(a_cy * scale)))


def _fix_cell_images(paragraphs: list, max_width_emu: int):
    """Procesa todos los párrafos: convierte anchors a inline y limita ancho.

    Args:
        paragraphs: Lista de elementos <w:p> XML.
        max_width_emu: Ancho máximo para imágenes en EMUs.
    """
    for p in paragraphs:
        _convert_anchors_to_inline(p)
        _clamp_image_width(p, max_width_emu)


def _get_column_widths_emu(table) -> list[int]:
    """Extrae los anchos de columna de <w:tblGrid> en EMUs.

    Args:
        table: Objeto Table de python-docx.

    Returns:
        Lista de anchos por columna en EMUs. Vacía si no hay tblGrid.
    """
    tblGrid = table._tbl.find(f"{{{NS_W}}}tblGrid")
    if tblGrid is None:
        return []

    widths = []
    for gc in tblGrid.findall(f"{{{NS_W}}}gridCol"):
        w_dxa_str = gc.get(f"{{{NS_W}}}w", "0")
        w_dxa = float(w_dxa_str)
        widths.append(int(w_dxa * 635))  # dxa → EMU
    return widths


def _build_cell_content(
    question: Question,
    max_image_width_emu: int = DEFAULT_MAX_IMAGE_WIDTH_EMU,
) -> list:
    """Construye la lista de párrafos XML para una celda a partir de una pregunta.

    Orden:
    1. Párrafos del enunciado (con número, formato e imágenes)
    2. Párrafos de cada alternativa (con letra, formato e imágenes)
    3. Párrafos vacíos de trailing (padding original)

    Después de ensamblar, procesa imágenes:
    - Convierte anchors flotantes a inline
    - Escala imágenes que excedan el ancho de celda

    Args:
        question: Pregunta randomizada con XML preservado.
        max_image_width_emu: Ancho máximo para imágenes en EMUs.

    Returns:
        Lista de elementos <w:p> lxml.
    """
    paragraphs = []

    # Enunciado — separar párrafos mixtos y reordenar: texto primero, imágenes siempre al final
    enunciado_text = []
    enunciado_images = []
    for p in question.enunciado_paragraphs:
        for split_p in _split_text_and_image_runs(p):
            if _paragraph_has_image(split_p):
                enunciado_images.append(split_p)
            else:
                enunciado_text.append(split_p)

    # Alternativas — detectar párrafos mixtos (texto de alternativa + imagen del enunciado).
    # Ocurre cuando en el .docx original la imagen fue insertada en el mismo párrafo que
    # una alternativa (ej: "c) 2 ciclos [imagen de onda]"). En ese caso se extrae la imagen
    # y se mueve al bloque del enunciado, quedando visible antes de las alternativas.
    clean_alt_paragraphs: list[list] = []  # una sub-lista por alternativa
    for alt in question.alternatives:
        alt_clean = []
        for p in alt.paragraphs:
            split = _split_text_and_image_runs(p)
            if len(split) == 2:
                # Párrafo mixto: el primer split es texto, el segundo es imagen.
                text_p, img_p = split
                enunciado_images.append(img_p)
                alt_clean.append(text_p)
            else:
                alt_clean.append(split[0])
        clean_alt_paragraphs.append(alt_clean)

    for p in enunciado_text:
        paragraphs.append(p)
    for p in enunciado_images:
        paragraphs.append(p)

    # Alternativas (ya limpias de imágenes extraídas)
    for alt_paras in clean_alt_paragraphs:
        for p in alt_paras:
            paragraphs.append(p)

    # Trailing empty paragraphs (preservar layout/spacing)
    for p in question.trailing_empty_paragraphs:
        paragraphs.append(p)

    # Procesar imágenes: anchor→inline + clamp width
    _fix_cell_images(paragraphs, max_image_width_emu)

    return paragraphs


def _update_form_label(doc: Document, form_number: int):
    """Busca y reemplaza el texto 'Forma N' en los párrafos superiores.

    Busca un párrafo que contenga la palabra "Forma" seguida de un número
    y lo reemplaza con el nuevo número de forma.

    Args:
        doc: Documento python-docx.
        form_number: Número de la forma (ej: 2, 3, 4).
    """
    form_pattern = re.compile(r"([Ff]orma\s+)\d+")

    for paragraph in doc.paragraphs:
        full_text = paragraph.text
        if form_pattern.search(full_text):
            for run in paragraph.runs:
                if form_pattern.search(run.text):
                    run.text = form_pattern.sub(rf"\g<1>{form_number}", run.text)
            break  # Solo reemplazar la primera ocurrencia


def _set_rows_cant_split(table):
    """Configura todas las filas de la tabla para que no se dividan entre páginas.

    Usa OxmlElement y qn de python-docx para agregar <w:cantSplit/> a
    <w:trPr> de cada fila, en lugar de manipulación XML cruda con lxml.

    Esto asegura que si una pregunta no cabe en el espacio restante de la
    página actual, Word mueve la fila completa al inicio de la siguiente página.

    Args:
        table: Objeto Table de python-docx.
    """
    for row in table.rows:
        tr = row._tr
        # Buscar o crear <w:trPr>
        trPr = tr.find(qn("w:trPr"))
        if trPr is None:
            trPr = OxmlElement("w:trPr")
            tr.insert(0, trPr)

        # Buscar o crear <w:cantSplit> — forzar val="1" para override
        cant_split = trPr.find(qn("w:cantSplit"))
        if cant_split is None:
            cant_split = OxmlElement("w:cantSplit")
            trPr.append(cant_split)
        cant_split.set(qn("w:val"), "1")


def build_docx(
    source_path: str,
    questions: list[Question],
    output_path: str,
    form_number: int,
    questions_table_index: int = 2,
    questions_per_row: int = 2,
) -> None:
    """Genera un .docx randomizado clonando el original y reemplazando preguntas.

    Args:
        source_path: Ruta al .docx original.
        questions: Lista de preguntas randomizadas (con XML preservado).
        output_path: Ruta donde guardar el nuevo .docx.
        form_number: Número de la forma (ej: 2, 3, 4) para actualizar el label.
        questions_table_index: Índice de la tabla de preguntas en el documento.
        questions_per_row: Número de preguntas por fila (columnas de la tabla).
    """
    # 1. Copiar el archivo original
    shutil.copy2(source_path, output_path)

    # 2. Abrir la copia y modificar
    doc = Document(output_path)

    # 3. Actualizar label "Forma X"
    _update_form_label(doc, form_number)

    # 4. Obtener la tabla de preguntas
    questions_table = doc.tables[questions_table_index]

    # 5. Configurar filas para que no se corten entre páginas
    _set_rows_cant_split(questions_table)

    # 6. Calcular ancho máximo de imagen por columna desde <w:tblGrid>
    col_widths_emu = _get_column_widths_emu(questions_table)

    # 7. Capturar contenido del body antes y después de la tabla de preguntas
    #    (texto de cierre, imágenes, sectPr) para preservarlo tras modificaciones
    body = doc.element.body
    table_element = questions_table._tbl
    # Registrar el índice exacto de la tabla en el body para poder reinsertarla
    # defensivamente si alguna operación la desanclara accidentalmente
    table_body_index = list(body).index(table_element)
    tail_elements = []
    found_table = False
    for child in list(body):
        if child is table_element:
            found_table = True
            continue
        if found_table:
            tail_elements.append(child)

    # 8. Reemplazar contenido de cada celda
    q_index = 0
    for row in questions_table.rows:
        for col_idx in range(questions_per_row):
            if q_index >= len(questions):
                break
            cell = row.cells[col_idx]

            # Determinar ancho máximo de imagen para esta columna
            if col_idx < len(col_widths_emu):
                max_img_width = col_widths_emu[col_idx] - CELL_PADDING_EMU
            else:
                max_img_width = DEFAULT_MAX_IMAGE_WIDTH_EMU

            # Limpiar la celda
            _clear_cell(cell)

            # Insertar los párrafos de la pregunta randomizada
            cell_content = _build_cell_content(questions[q_index], max_img_width)
            _add_paragraphs_to_cell(cell, cell_content)

            q_index += 1

    # 9. Re-asegurar que la tabla y el contenido posterior siguen en el body
    #    en el orden correcto. La tabla ya fue modificada in-place, así que
    #    solo actuamos si fue removida accidentalmente del body.
    body_children = list(body)
    if table_element not in body_children:
        # La tabla fue removida accidentalmente — reinsertarla en su posición original
        if table_body_index <= len(body_children):
            body.insert(table_body_index, table_element)
        else:
            body.append(table_element)

    # 10. Guardar
    doc.save(output_path)
