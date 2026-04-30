"""Parser de archivos .docx con formato de pruebas de alternativas en tabla.

Estructura esperada del documento:
- Tablas iniciales con datos del alumno e instrucciones
- Una tabla principal (normalmente la última/más grande) con preguntas
- Cada celda de la tabla de preguntas contiene:
  - Enunciado: párrafo(s) que empiezan con "N-" o "N."
  - Alternativas: párrafos que empiezan con "x)" donde x es letra
  - Párrafos vacíos de padding (al final)

El parser extrae directamente los elementos XML <w:p> para preservar
todo el formato (negrita, cursiva, imágenes, colores, fuentes).
"""

import re
from copy import deepcopy
from typing import List, Optional

from docx import Document

from src.models import Alternative, Exam, Question

# Namespace de Word XML
NSMAP = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

# Frases que marcan el inicio del ítem de alternativas (case-insensitive, coincidencia parcial)
_ALTERNATIVE_MARKERS = (
    "marque la alternativa",
    "marca la alternativa",
    "seleccione la alternativa",
    "selecciona la alternativa",
)

# Tag local para párrafo y tabla en el body del documento
_W_P_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
_W_TBL_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl"

# Regex para detectar inicio de pregunta: "N-" o "N." donde N es uno o más dígitos
QUESTION_PATTERN = re.compile(r"^(\d+)\s*[-–—.]\s*")

# Regex para detectar inicio de alternativa: "x)" donde x es una letra
ALTERNATIVE_PATTERN = re.compile(r"^([a-zA-Z])\)\s*")


def _get_paragraph_text(p_element) -> str:
    """Extrae texto plano de un elemento <w:p> XML.

    Args:
        p_element: Elemento lxml representando un párrafo de Word.

    Returns:
        Texto concatenado de todos los runs del párrafo.
    """
    runs = p_element.findall(".//w:r/w:t", NSMAP)
    return "".join(r.text or "" for r in runs).strip()


def _paragraph_has_image(p_element) -> bool:
    """Detecta si un párrafo contiene imágenes embebidas.

    Busca elementos <w:drawing> (imágenes modernas) y <w:pict> (legacy)
    dentro del párrafo. También hace fallback iterando todos los
    descendientes buscando tags relacionados a imágenes.

    Args:
        p_element: Elemento lxml representando un párrafo de Word.

    Returns:
        True si el párrafo contiene al menos una imagen.
    """
    # <w:drawing> — imágenes modernas (Word 2007+)
    if p_element.findall(".//w:drawing", NSMAP):
        return True
    # <w:pict> — imágenes legacy
    if p_element.findall(".//w:pict", NSMAP):
        return True
    # Fallback: buscar en cualquier namespace
    for elem in p_element.iter():
        tag = elem.tag if isinstance(elem.tag, str) else ""
        if "drawing" in tag or "blip" in tag or "imagedata" in tag:
            return True
    return False


def _find_questions_table_smart(doc: Document):
    """Encuentra la tabla de preguntas usando detección inteligente por marcador.

    Estrategia principal: recorre el body del documento en orden buscando un
    párrafo que contenga alguna de las frases marcadoras definidas en
    _ALTERNATIVE_MARKERS (case-insensitive). La primera tabla que aparezca
    *después* de ese párrafo es considerada la tabla de preguntas.

    Fallback: si no se encuentra el marcador, usa la heurística de la tabla con
    más filas (comportamiento idéntico a la función original).

    Args:
        doc: Documento python-docx.

    Returns:
        Tupla (índice_tabla_en_doc.tables, objeto_tabla).

    Raises:
        ValueError: Si el documento no contiene tablas.
    """
    if not doc.tables:
        raise ValueError("El documento no contiene tablas.")

    body = doc.element.body
    found_marker = False

    for child in body:
        tag = child.tag if isinstance(child.tag, str) else ""

        # Es un párrafo — buscar el marcador si aún no lo hemos encontrado
        if tag == _W_P_TAG and not found_marker:
            text = _get_paragraph_text(child).lower()
            if any(marker in text for marker in _ALTERNATIVE_MARKERS):
                found_marker = True

        # Es una tabla — si ya encontramos el marcador, esta es la tabla de preguntas
        elif tag == _W_TBL_TAG and found_marker:
            # Buscar el índice de esta tabla dentro de doc.tables comparando por identidad
            for idx, table in enumerate(doc.tables):
                if table._tbl is child:
                    return idx, table

    # Fallback: heurística de tabla con más filas
    markers_display = ", ".join(f'"{m}"' for m in _ALTERNATIVE_MARKERS)
    print(
        f"⚠️  No se encontró ninguna frase marcadora ({markers_display}). "
        "Usando heurística de tabla más grande."
    )
    max_rows = 0
    best_idx = 0
    for i, table in enumerate(doc.tables):
        if len(table.rows) > max_rows:
            max_rows = len(table.rows)
            best_idx = i

    return best_idx, doc.tables[best_idx]


def _parse_cell(cell) -> Optional[Question]:
    """Parsea una celda de la tabla de preguntas.

    Extrae el enunciado, alternativas y párrafos vacíos de trailing.
    Preserva los elementos XML originales (deep copy) para mantener formato.

    Args:
        cell: Objeto Cell de python-docx.

    Returns:
        Un objeto Question si la celda contiene una pregunta, None si está vacía.
    """
    # Acceder a los elementos XML <w:p> dentro del <w:tc>
    tc_element = cell._tc
    p_elements = tc_element.findall("w:p", NSMAP)

    if not p_elements:
        return None

    # Verificar que la celda contiene una pregunta
    first_text = _get_paragraph_text(p_elements[0])
    q_match = QUESTION_PATTERN.match(first_text)
    if not q_match:
        return None

    question_number = int(q_match.group(1))

    # Clasificar párrafos: enunciado, alternativas, trailing empty
    enunciado_paragraphs = []
    alternatives = []
    trailing_empty = []
    current_alt_paragraphs = []
    current_alt_letter = None
    in_alternatives = False

    for p_elem in p_elements:
        text = _get_paragraph_text(p_elem)
        has_image = _paragraph_has_image(p_elem)
        is_empty = not text and not has_image
        alt_match = ALTERNATIVE_PATTERN.match(text)

        if alt_match:
            # Si había una alternativa en curso, guardarla
            if current_alt_letter is not None:
                alternatives.append(
                    Alternative(
                        letter=current_alt_letter,
                        paragraphs=[deepcopy(pp) for pp in current_alt_paragraphs],
                    )
                )

            in_alternatives = True
            current_alt_letter = alt_match.group(1).lower()
            current_alt_paragraphs = [p_elem]
            # Reset trailing (previous empty paragraphs were part of content)
            if trailing_empty:
                # Were these actually between alternatives? Add to previous alt
                if alternatives:
                    alternatives[-1].paragraphs.extend(
                        [deepcopy(pp) for pp in trailing_empty]
                    )
                trailing_empty = []

        elif is_empty and in_alternatives:
            # Párrafo verdaderamente vacío después de las alternativas — es trailing
            trailing_empty.append(p_elem)

        elif is_empty and not in_alternatives:
            # Párrafo vacío en el enunciado — separador de spacing
            enunciado_paragraphs.append(p_elem)

        else:
            # Párrafo con contenido (texto o imagen)
            if not in_alternatives:
                enunciado_paragraphs.append(p_elem)
            else:
                # Contenido (texto o imagen) después de alternativas —
                # parte de la alternativa actual
                current_alt_paragraphs.append(p_elem)
                trailing_empty = []

    # Guardar la última alternativa
    if current_alt_letter is not None:
        alternatives.append(
            Alternative(
                letter=current_alt_letter,
                paragraphs=[deepcopy(pp) for pp in current_alt_paragraphs],
            )
        )

    if not alternatives:
        return None

    return Question(
        number=question_number,
        enunciado_paragraphs=[deepcopy(p) for p in enunciado_paragraphs],
        alternatives=alternatives,
        trailing_empty_paragraphs=[deepcopy(p) for p in trailing_empty],
    )


def parse_docx(file_path: str) -> Exam:
    """Parsea un archivo .docx de prueba de alternativas con preguntas en tabla.

    Args:
        file_path: Ruta al archivo .docx.

    Returns:
        Un objeto Exam con las preguntas extraídas, preservando elementos XML.
    """
    doc = Document(file_path)

    table_idx, questions_table = _find_questions_table_smart(doc)

    questions: List[Question] = []

    for row in questions_table.rows:
        for cell in row.cells:
            question = _parse_cell(cell)
            if question is not None:
                questions.append(question)

    # Determinar preguntas por fila
    cols_count = len(questions_table.columns) if questions_table.columns else 2

    # Recopilar elementos del body antes y después de la tabla de preguntas
    questions_tbl_element = questions_table._tbl
    pre_elements = []
    post_elements = []
    found_table = False

    for child in doc.element.body:
        if child is questions_tbl_element:
            found_table = True
            continue  # excluir la tabla misma
        if found_table:
            post_elements.append(deepcopy(child))
        else:
            pre_elements.append(deepcopy(child))

    return Exam(
        source_path=file_path,
        questions=questions,
        questions_table_index=table_idx,
        questions_per_row=cols_count,
        pre_table_body_elements=pre_elements,
        post_table_body_elements=post_elements,
    )
