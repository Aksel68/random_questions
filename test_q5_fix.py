#!/usr/bin/env python3
"""Test rápido para verificar que la imagen de la pregunta 5 queda en el enunciado.

Uso:
    python test_q5_fix.py --drive "Guías/Ondas 2026. Ev. 1 F2.docx"
    python test_q5_fix.py --local /ruta/al/archivo.docx
"""

import argparse
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from copy import deepcopy

from docx import Document
from lxml import etree

from src.doc_builder import (
    CELL_PADDING_EMU,
    DEFAULT_MAX_IMAGE_WIDTH_EMU,
    _build_cell_content,
    _get_column_widths_emu,
)
from src.parser import parse_docx

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NSMAP = {"w": NS_W}


def _has_image(p_element) -> bool:
    if p_element.findall(f".//{{{NS_W}}}drawing"):
        return True
    if p_element.findall(f".//{{{NS_MC}}}AlternateContent"):
        return True
    if p_element.findall(".//w:pict", NSMAP):
        return True
    for elem in p_element.iter():
        tag = elem.tag if isinstance(elem.tag, str) else ""
        if "drawing" in tag or "blip" in tag or "imagedata" in tag:
            return True
    return False


def _get_text(p_element) -> str:
    runs = p_element.findall(".//w:r/w:t", NSMAP)
    return "".join(r.text or "" for r in runs).strip()


def run_test(file_path: str, question_number: int = 5):
    print(f"\n{'=' * 60}")
    print(f"  TEST: imagen de pregunta {question_number} va al enunciado")
    print(f"  Archivo: {file_path}")
    print(f"{'=' * 60}\n")

    # 1. Parsear
    exam = parse_docx(file_path)
    question = next((q for q in exam.questions if q.number == question_number), None)

    if question is None:
        print(f"❌  No se encontró la pregunta {question_number}.")
        sys.exit(1)

    print(f"Pregunta {question_number} encontrada.")
    print(f"  Párrafos enunciado : {len(question.enunciado_paragraphs)}")
    print(f"  Alternativas       : {len(question.alternatives)}")

    # Mostrar estado PRE-fix (lo que viene del parser)
    print("\n--- Estado del parser (antes de _build_cell_content) ---")
    img_in_enunciado_before = sum(
        1 for p in question.enunciado_paragraphs if _has_image(p)
    )
    img_in_alts_before = sum(
        1 for alt in question.alternatives for p in alt.paragraphs if _has_image(p)
    )
    print(f"  Imágenes en enunciado : {img_in_enunciado_before}")
    print(f"  Imágenes en alternativas: {img_in_alts_before}")

    for i, alt in enumerate(question.alternatives):
        for p in alt.paragraphs:
            if _has_image(p):
                txt = _get_text(p)
                print(
                    f"    ⚠️  Alt [{alt.letter}] tiene imagen — texto: {repr(txt[:60])}"
                )

    # 2. Obtener doc para calcular ancho de columna
    doc = Document(file_path)
    questions_table = doc.tables[exam.questions_table_index]
    col_widths_emu = _get_column_widths_emu(questions_table)
    max_img_width = (
        col_widths_emu[0] - CELL_PADDING_EMU
        if col_widths_emu
        else DEFAULT_MAX_IMAGE_WIDTH_EMU
    )

    # 3. Construir contenido de celda (aplica el fix)
    paragraphs = _build_cell_content(question, max_img_width)

    # 4. Analizar resultado
    print("\n--- Resultado de _build_cell_content (después del fix) ---")
    print(f"  Total párrafos generados: {len(paragraphs)}")

    import re

    ALT_PATTERN = re.compile(r"^([a-zA-Z])\)\s*")

    in_alts = False
    images_before_alts = 0
    images_after_alts = 0
    alt_with_image = []

    for i, p in enumerate(paragraphs):
        txt = _get_text(p)
        has_img = _has_image(p)
        is_alt = bool(ALT_PATTERN.match(txt))

        if is_alt:
            in_alts = True

        if has_img:
            if not in_alts:
                images_before_alts += 1
            else:
                images_after_alts += 1
                alt_with_image.append((i, txt[:50]))

        status = "ALT" if in_alts else "ENE"
        img_flag = " 🖼" if has_img else ""
        print(f"  [{i:02d}] [{status}] {repr(txt[:55])}{img_flag}")

    print()
    if images_before_alts > 0 and images_after_alts == 0:
        print(
            f"✅  PASS — La imagen quedó ANTES de las alternativas (en el enunciado)."
        )
    elif images_after_alts > 0:
        print(f"❌  FAIL — La imagen todavía está DENTRO de una alternativa:")
        for idx, txt in alt_with_image:
            print(f"     párrafo [{idx:02d}]: {repr(txt)}")
    elif images_before_alts == 0 and images_after_alts == 0:
        print("⚠️  No se encontró ninguna imagen en el contenido generado.")

    return images_before_alts > 0 and images_after_alts == 0


def main():
    parser = argparse.ArgumentParser(
        description="Verifica que la imagen de una pregunta quede en el enunciado, no en las alternativas."
    )
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--local", metavar="PATH", help="Ruta local al archivo .docx")
    group.add_argument("--drive", metavar="DRIVE_PATH", help="Ruta en Google Drive")
    parser.add_argument(
        "--question", type=int, default=5, help="Número de pregunta (default: 5)"
    )
    args = parser.parse_args()

    if args.local:
        file_path = args.local
    else:
        print("⬇️  Descargando desde Google Drive...")
        from src.drive_client import download_docx

        file_path = download_docx(args.drive)
        print(f"   Descargado en: {file_path}\n")

    passed = run_test(file_path, args.question)
    sys.exit(0 if passed else 1)


if __name__ == "__main__":
    main()
