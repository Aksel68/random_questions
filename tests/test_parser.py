import os
import tempfile

from docx import Document

from src.models import Exam
from src.parser import parse_docx


def _create_table_exam_docx(path: str, num_questions: int = 6):
    """Crea un .docx con estructura de tabla similar al documento real.

    - Tabla 0 (1x2): Header (simulado)
    - Tabla 1 (Nx2): Preguntas en pares (2 por fila)
    """
    doc = Document()
    doc.add_paragraph("Prueba de Física")

    # Tabla de datos (header)
    header_table = doc.add_table(rows=1, cols=2)
    header_table.cell(0, 0).text = "NOMBRE"
    header_table.cell(0, 1).text = ""

    doc.add_paragraph("Ítem I. Marque la alternativa correcta.")

    # Tabla de preguntas (2 columnas)
    rows_needed = (num_questions + 1) // 2
    q_table = doc.add_table(rows=rows_needed, cols=2)

    questions_data = [
        (
            "¿Cuál es la velocidad de la luz?",
            ["3x10^8 m/s", "3x10^6 m/s", "3x10^10 m/s", "3x10^4 m/s"],
        ),
        (
            "¿Qué tipo de onda es el sonido?",
            ["Transversal", "Longitudinal", "Electromagnética", "Estacionaria"],
        ),
        ("¿Cuál es la unidad de frecuencia?", ["Metro", "Segundo", "Hertz", "Newton"]),
        ("¿Qué transmiten las ondas?", ["Materia", "Fuerza", "Energía", "Calor"]),
        (
            "¿Qué es un pulso?",
            ["Perturbación aislada", "Onda periódica", "Onda continua", "Vibración"],
        ),
        (
            "¿Qué necesitan las ondas mecánicas?",
            ["Vacío", "Luz", "Medio material", "Electricidad"],
        ),
    ]

    for i in range(min(num_questions, len(questions_data))):
        row_idx = i // 2
        col_idx = i % 2
        cell = q_table.cell(row_idx, col_idx)

        q_text, alts = questions_data[i]
        # Limpiar celda (viene con un párrafo vacío por defecto)
        cell.paragraphs[0].text = f"{i + 1}- {q_text}"

        for j, alt_text in enumerate(alts):
            letter = chr(ord("a") + j)
            cell.add_paragraph(f"{letter}) {alt_text}")

    doc.save(path)


def test_parse_docx_extracts_questions_from_table():
    """Verifica que el parser extrae preguntas de la tabla correctamente."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name
    try:
        _create_table_exam_docx(path, num_questions=6)
        exam = parse_docx(path)

        assert isinstance(exam, Exam)
        assert len(exam.questions) == 6
        assert exam.questions_per_row == 2

        # Verificar primera pregunta
        q1 = exam.questions[0]
        assert q1.number == 1
        assert "velocidad de la luz" in q1.text.lower()
        assert len(q1.alternatives) == 4
        assert q1.alternatives[0].letter == "a"
    finally:
        os.unlink(path)


def test_parse_docx_preserves_formatting():
    """Verifica que el parser preserva los elementos XML con formato."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name
    try:
        doc = Document()
        q_table = doc.add_table(rows=1, cols=1)
        cell = q_table.cell(0, 0)

        # Pregunta con negrita
        p = cell.paragraphs[0]
        run1 = p.add_run("1- ¿Qué es la ")
        run2 = p.add_run("amplitud")
        run2.bold = True
        run3 = p.add_run("?")

        cell.add_paragraph("a) Distancia máxima")
        cell.add_paragraph("b) Frecuencia")
        cell.add_paragraph("c) Período")
        cell.add_paragraph("d) Velocidad")

        doc.save(path)

        exam = parse_docx(path)
        assert len(exam.questions) == 1
        q = exam.questions[0]
        assert q.number == 1
        assert len(q.alternatives) == 4

        # Verificar que el XML del enunciado tiene runs con formato
        nsmap = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        runs = q.enunciado_paragraphs[0].findall(".//w:r", nsmap)
        # Debe haber al menos 3 runs (texto + bold + texto)
        assert len(runs) >= 3
    finally:
        os.unlink(path)


def test_parse_docx_handles_dot_format():
    """Verifica que el parser detecta preguntas con formato 'N.' además de 'N-'."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name
    try:
        doc = Document()
        q_table = doc.add_table(rows=1, cols=1)
        cell = q_table.cell(0, 0)
        cell.paragraphs[0].text = "33. De los siguientes, ¿cuál NO es una onda?"
        cell.add_paragraph("a) Luz")
        cell.add_paragraph("b) Sonido")
        cell.add_paragraph("c) Vuelo de abeja")
        cell.add_paragraph("d) Terremoto")

        doc.save(path)

        exam = parse_docx(path)
        assert len(exam.questions) == 1
        assert exam.questions[0].number == 33
    finally:
        os.unlink(path)


def test_parse_docx_handles_multiline_enunciado():
    """Verifica que el enunciado multilinea se captura completo."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name
    try:
        doc = Document()
        q_table = doc.add_table(rows=1, cols=1)
        cell = q_table.cell(0, 0)
        cell.paragraphs[0].text = '3-"Corresponde a la posición de reposo en el'
        cell.add_paragraph("que se encuentra un cuerpo antes de")
        cell.add_paragraph('empezar a vibrar". Esta definición corresponde a:')
        cell.add_paragraph("a) Amplitud")
        cell.add_paragraph("b) Longitud de onda")
        cell.add_paragraph("c) Valle")
        cell.add_paragraph("d) Punto de equilibrio")

        doc.save(path)

        exam = parse_docx(path)
        assert len(exam.questions) == 1
        q = exam.questions[0]
        assert q.number == 3
        assert len(q.enunciado_paragraphs) == 3
        assert len(q.alternatives) == 4
    finally:
        os.unlink(path)


def test_parse_finds_correct_table():
    """Verifica que el parser identifica la tabla de preguntas (la más grande)."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name
    try:
        doc = Document()
        # Tabla pequeña (header)
        t1 = doc.add_table(rows=2, cols=4)
        t1.cell(0, 0).text = "NOMBRE"

        # Tabla pequeña (instrucciones)
        t2 = doc.add_table(rows=1, cols=1)
        t2.cell(0, 0).text = "INSTRUCCIONES"

        # Tabla grande (preguntas) — esta debe ser detectada
        t3 = doc.add_table(rows=5, cols=2)
        for i in range(10):
            row_idx = i // 2
            col_idx = i % 2
            cell = t3.cell(row_idx, col_idx)
            cell.paragraphs[0].text = f"{i + 1}- Pregunta {i + 1}"
            cell.add_paragraph("a) Opción A")
            cell.add_paragraph("b) Opción B")
            cell.add_paragraph("c) Opción C")
            cell.add_paragraph("d) Opción D")

        doc.save(path)

        exam = parse_docx(path)
        assert exam.questions_table_index == 2  # Tercera tabla (index 2)
        assert len(exam.questions) == 10
    finally:
        os.unlink(path)
