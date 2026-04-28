# tests/test_models.py
"""Tests de los modelos v2 — XML-based."""

import os
import tempfile

from docx import Document

from src.models import Alternative, Exam, Question
from src.parser import parse_docx


def _create_simple_exam() -> Exam:
    """Helper: crea un exam simple parseando un docx temporal."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name

    doc = Document()
    qt = doc.add_table(rows=1, cols=1)
    cell = qt.cell(0, 0)
    p = cell.paragraphs[0]
    p.add_run("1- ")
    bold_run = p.add_run("Pregunta")
    bold_run.bold = True
    p.add_run(" de prueba")
    cell.add_paragraph("a) Opción A")
    cell.add_paragraph("b) Opción B")
    doc.save(path)

    exam = parse_docx(path)
    os.unlink(path)
    return exam


def test_question_text_property():
    """El .text debe devolver el texto plano del enunciado."""
    exam = _create_simple_exam()
    q = exam.questions[0]
    assert "Pregunta" in q.text
    assert "prueba" in q.text


def test_alternative_text_property():
    """El .text debe devolver el texto plano de la alternativa."""
    exam = _create_simple_exam()
    q = exam.questions[0]
    assert q.alternatives[0].text.strip()  # No vacío
    assert "Opción A" in q.alternatives[0].text or "a)" in q.alternatives[0].text


def test_question_clone_is_independent():
    """Clonar una pregunta no debe afectar el original."""
    exam = _create_simple_exam()
    q = exam.questions[0]
    original_text = q.text
    cloned = q.clone()

    # Modificar el clon no afecta al original
    cloned.number = 99
    assert q.number == 1
    assert q.text == original_text


def test_alternative_clone_is_independent():
    """Clonar una alternativa no debe afectar el original."""
    exam = _create_simple_exam()
    alt = exam.questions[0].alternatives[0]
    original_letter = alt.letter
    cloned = alt.clone()

    cloned.letter = "z"
    assert alt.letter == original_letter


def test_alternative_original_letter_defaults_to_letter():
    """original_letter debe ser igual a letter si no se especifica."""
    from lxml import etree

    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    p = etree.SubElement(etree.Element("root"), f"{{{ns}}}p")
    r = etree.SubElement(p, f"{{{ns}}}r")
    t = etree.SubElement(r, f"{{{ns}}}t")
    t.text = "a) Opción A"

    alt = Alternative(letter="a", paragraphs=[p])
    assert alt.original_letter == "a"


def test_alternative_original_letter_preserved_on_clone():
    """clone() debe preservar original_letter."""
    from lxml import etree

    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    p = etree.SubElement(etree.Element("root"), f"{{{ns}}}p")

    alt = Alternative(letter="c", paragraphs=[p], original_letter="a")
    cloned = alt.clone()
    assert cloned.original_letter == "a"
    assert cloned.letter == "c"


def test_question_original_number_defaults_to_number():
    """original_number debe ser igual a number si no se especifica."""
    q = Question(number=5, enunciado_paragraphs=[])
    assert q.original_number == 5


def test_question_original_number_preserved_on_clone():
    """clone() debe preservar original_number."""
    q = Question(number=12, enunciado_paragraphs=[], original_number=3)
    cloned = q.clone()
    assert cloned.original_number == 3
    assert cloned.number == 12
