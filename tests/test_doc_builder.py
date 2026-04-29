import os
import tempfile

from docx import Document
from docx.oxml.ns import qn

from src.doc_builder import build_docx
from src.models import Exam
from src.parser import parse_docx
from src.randomizer import generate_forms


def _create_source_exam(path: str):
    """Crea un documento fuente con tabla de preguntas y formato."""
    from docx.shared import Pt

    doc = Document()

    # Título con formato
    p = doc.add_paragraph()
    run = p.add_run("Prueba de Física - Ondas")
    run.bold = True
    run.font.size = Pt(14)

    # Label de forma
    p2 = doc.add_paragraph()
    run2 = p2.add_run("Forma 1")
    run2.bold = True

    # Tabla header
    ht = doc.add_table(rows=1, cols=2)
    ht.cell(0, 0).text = "NOMBRE"
    ht.cell(0, 1).text = ""

    # Tabla instrucciones
    it = doc.add_table(rows=1, cols=1)
    it.cell(0, 0).text = "INSTRUCCIONES: Lea con atención"

    # Tabla preguntas (3 filas x 2 cols = 6 preguntas)
    qt = doc.add_table(rows=3, cols=2)
    for i in range(6):
        row_idx = i // 2
        col_idx = i % 2
        cell = qt.cell(row_idx, col_idx)
        # Pregunta con negrita parcial
        p = cell.paragraphs[0]
        p.add_run(f"{i + 1}- ¿Qué es la ")
        bold_run = p.add_run("onda")
        bold_run.bold = True
        p.add_run(f" número {i + 1}?")

        cell.add_paragraph(f"a) Opción A")
        cell.add_paragraph(f"b) Opción B")
        cell.add_paragraph(f"c) Opción C")
        cell.add_paragraph(f"d) Opción D")

    doc.save(path)


def test_build_docx_creates_file():
    """Verifica que build_docx genera un archivo con contenido."""
    with tempfile.TemporaryDirectory() as tmpdir:
        source = os.path.join(tmpdir, "original.docx")
        output = os.path.join(tmpdir, "forma_a.docx")
        _create_source_exam(source)

        exam = parse_docx(source)
        forms = generate_forms(exam, n=1)

        build_docx(
            source_path=source,
            questions=forms[0],
            output_path=output,
            form_number=2,
            questions_table_index=exam.questions_table_index,
            questions_per_row=exam.questions_per_row,
        )

        assert os.path.exists(output)
        assert os.path.getsize(output) > 0


def test_build_docx_preserves_header_tables():
    """Las tablas de header/instrucciones deben permanecer intactas."""
    with tempfile.TemporaryDirectory() as tmpdir:
        source = os.path.join(tmpdir, "original.docx")
        output = os.path.join(tmpdir, "forma_a.docx")
        _create_source_exam(source)

        exam = parse_docx(source)
        forms = generate_forms(exam, n=1)

        build_docx(
            source_path=source,
            questions=forms[0],
            output_path=output,
            form_number=2,
            questions_table_index=exam.questions_table_index,
            questions_per_row=exam.questions_per_row,
        )

        out_doc = Document(output)
        # Tabla 0 debe tener "NOMBRE"
        assert "NOMBRE" in out_doc.tables[0].cell(0, 0).text
        # Tabla 1 debe tener "INSTRUCCIONES"
        assert "INSTRUCCIONES" in out_doc.tables[1].cell(0, 0).text


def test_build_docx_updates_form_label():
    """El label 'Forma X' debe actualizarse al nuevo sufijo."""
    with tempfile.TemporaryDirectory() as tmpdir:
        source = os.path.join(tmpdir, "original.docx")
        output = os.path.join(tmpdir, "forma_b.docx")
        _create_source_exam(source)

        exam = parse_docx(source)
        forms = generate_forms(exam, n=1)

        build_docx(
            source_path=source,
            questions=forms[0],
            output_path=output,
            form_number=2,
            questions_table_index=exam.questions_table_index,
            questions_per_row=exam.questions_per_row,
        )

        out_doc = Document(output)
        full_text = " ".join(p.text for p in out_doc.paragraphs)
        assert "Forma 2" in full_text
        assert "Forma 1" not in full_text


def test_build_docx_contains_all_questions():
    """El documento generado debe contener todas las preguntas."""
    with tempfile.TemporaryDirectory() as tmpdir:
        source = os.path.join(tmpdir, "original.docx")
        output = os.path.join(tmpdir, "forma_a.docx")
        _create_source_exam(source)

        exam = parse_docx(source)
        forms = generate_forms(exam, n=1)

        build_docx(
            source_path=source,
            questions=forms[0],
            output_path=output,
            form_number=2,
            questions_table_index=exam.questions_table_index,
            questions_per_row=exam.questions_per_row,
        )

        # Re-parsear para verificar
        out_exam = parse_docx(output)
        assert len(out_exam.questions) == 6
        for q in out_exam.questions:
            assert len(q.alternatives) == 4


def test_build_docx_preserves_bold_formatting():
    """El formato negrita del enunciado debe preservarse."""
    with tempfile.TemporaryDirectory() as tmpdir:
        source = os.path.join(tmpdir, "original.docx")
        output = os.path.join(tmpdir, "forma_a.docx")
        _create_source_exam(source)

        exam = parse_docx(source)
        # No randomizar — usar preguntas en orden original para verificar formato
        questions = [q.clone() for q in exam.questions]

        build_docx(
            source_path=source,
            questions=questions,
            output_path=output,
            form_number=2,
            questions_table_index=exam.questions_table_index,
            questions_per_row=exam.questions_per_row,
        )

        out_doc = Document(output)
        qt = out_doc.tables[exam.questions_table_index]
        first_cell = qt.cell(0, 0)
        first_para = first_cell.paragraphs[0]

        # Verificar que hay un run con bold
        has_bold = any(run.bold for run in first_para.runs)
        assert has_bold, "Se perdió el formato negrita en la primera celda"


def test_build_docx_sets_cant_split_on_rows():
    """Las filas de la tabla de preguntas deben tener cantSplit para evitar cortes."""
    with tempfile.TemporaryDirectory() as tmpdir:
        source = os.path.join(tmpdir, "original.docx")
        output = os.path.join(tmpdir, "output.docx")
        _create_source_exam(source)

        exam = parse_docx(source)
        forms = generate_forms(exam, n=1)

        build_docx(
            source_path=source,
            questions=forms[0],
            output_path=output,
            form_number=2,
            questions_table_index=exam.questions_table_index,
            questions_per_row=exam.questions_per_row,
        )

        out_doc = Document(output)
        qt = out_doc.tables[exam.questions_table_index]

        for row in qt.rows:
            tr = row._tr
            trPr = tr.find(qn("w:trPr"))
            assert trPr is not None, "La fila debe tener <w:trPr>"
            cant_split = trPr.find(qn("w:cantSplit"))
            assert cant_split is not None, "La fila debe tener <w:cantSplit/>"
