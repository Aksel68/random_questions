"""Tests para el módulo de generación de pautas de respuestas."""

import os
import tempfile

import openpyxl
import pytest
from lxml import etree
from src.answer_key import (
    build_answer_key_xlsx,
    generate_answer_key,
    parse_answer_key_csv,
)
from src.main import _complete_answer_key_or_none, _parse_local_xlsx_answer_key

from src.models import Alternative, Question

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _make_p(text: str):
    p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
    r = etree.SubElement(p, f"{{{NS_W}}}r")
    t = etree.SubElement(r, f"{{{NS_W}}}t")
    t.text = text
    return p


def _make_question(
    number: int, original_number: int, alt_mapping: list[tuple[str, str]]
) -> Question:
    """Crea una Question con alternativas que tienen letter y original_letter.

    Args:
        number: Número actual (en la forma randomizada).
        original_number: Número original (en el examen original).
        alt_mapping: Lista de (current_letter, original_letter).
    """
    alts = []
    for current, original in alt_mapping:
        alts.append(
            Alternative(
                letter=current,
                paragraphs=[_make_p(f"{current}) texto")],
                original_letter=original,
            )
        )
    return Question(
        number=number,
        enunciado_paragraphs=[_make_p(f"{number}- Pregunta")],
        alternatives=alts,
        original_number=original_number,
    )


class TestParseAnswerKeyCsv:
    def test_parses_simple_csv(self):
        csv_content = "Pregunta,Respuesta\n1,a\n2,b\n3,c\n4,d\n"
        result = parse_answer_key_csv(csv_content)
        assert result == {1: "a", 2: "b", 3: "c", 4: "d"}

    def test_parses_36_questions(self):
        lines = ["Pregunta,Respuesta"]
        letters = "abcd"
        for i in range(1, 37):
            lines.append(f"{i},{letters[(i - 1) % 4]}")
        csv_content = "\n".join(lines) + "\n"
        result = parse_answer_key_csv(csv_content)
        assert len(result) == 36
        assert result[1] == "a"
        assert result[36] == "d"

    def test_strips_whitespace(self):
        csv_content = "Pregunta,Respuesta\n 1 , a \n 2 , b \n"
        result = parse_answer_key_csv(csv_content)
        assert result == {1: "a", 2: "b"}

    def test_empty_csv_returns_empty(self):
        csv_content = "Pregunta,Respuesta\n"
        result = parse_answer_key_csv(csv_content)
        assert result == {}

    def test_lowercase_normalizes_uppercase(self):
        csv_content = "Pregunta,Respuesta\n1,A\n2,B\n"
        result = parse_answer_key_csv(csv_content)
        assert result[1] == "a"
        assert result[2] == "b"


class TestGenerateAnswerKey:
    def test_simple_mapping(self):
        """Q original 1 tiene respuesta 'a'. orig-a ahora es 'c' → respuesta = 'c'."""
        original_key = {1: "a", 2: "b"}

        q1 = _make_question(
            1,
            original_number=1,
            alt_mapping=[
                ("a", "d"),
                ("b", "b"),
                ("c", "a"),
                ("d", "c"),
            ],
        )
        q2 = _make_question(
            2,
            original_number=2,
            alt_mapping=[
                ("a", "b"),
                ("b", "c"),
                ("c", "d"),
                ("d", "a"),
            ],
        )

        result = generate_answer_key([q1, q2], original_key)
        assert result == {1: "c", 2: "a"}

    def test_preserves_all_questions(self):
        original_key = {1: "a", 2: "b", 3: "c"}
        questions = [
            _make_question(1, 3, [("a", "c"), ("b", "a"), ("c", "d"), ("d", "b")]),
            _make_question(2, 1, [("a", "b"), ("b", "a"), ("c", "d"), ("d", "c")]),
            _make_question(3, 2, [("a", "d"), ("b", "c"), ("c", "b"), ("d", "a")]),
        ]
        result = generate_answer_key(questions, original_key)
        # Q1 nueva (orig Q3, resp "c"): ("a","c") → current "a"
        assert result[1] == "a"
        # Q2 nueva (orig Q1, resp "a"): ("b","a") → current "b"
        assert result[2] == "b"
        # Q3 nueva (orig Q2, resp "b"): ("c","b") → current "c"
        assert result[3] == "c"
        assert len(result) == 3

    def test_missing_original_raises(self):
        original_key = {1: "a"}
        q = _make_question(
            1,
            original_number=99,
            alt_mapping=[
                ("a", "a"),
                ("b", "b"),
                ("c", "c"),
                ("d", "d"),
            ],
        )
        with pytest.raises(KeyError, match="pregunta original 99"):
            generate_answer_key([q], original_key)


class TestAnswerKeyCoverage:
    def test_incomplete_key_is_skipped_with_missing_question_numbers(self, capsys):
        questions = [
            _make_question(1, 1, [("a", "a")]),
            _make_question(2, 2, [("a", "a")]),
        ]

        result = _complete_answer_key_or_none({1: "a"}, questions)

        assert result is None
        assert "preguntas originales: 2" in capsys.readouterr().out

    def test_empty_key_is_skipped(self, capsys):
        question = _make_question(1, 1, [("a", "a")])

        result = _complete_answer_key_or_none({}, [question])

        assert result is None
        assert "no contiene respuestas válidas" in capsys.readouterr().out

    def test_complete_key_is_preserved(self):
        question = _make_question(1, 1, [("a", "a")])
        answer_key = {1: "a"}

        assert _complete_answer_key_or_none(answer_key, [question]) == answer_key


class TestParseLocalXlsxAnswerKey:
    def test_parses_integer_question_numbers_stored_as_floats(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Pregunta", "Respuesta"])
            sheet.append([1.0, "b"])
            sheet.append([17.0, "c"])
            workbook.save(path)

            assert _parse_local_xlsx_answer_key(path) == {1: "b", 17: "c"}
        finally:
            os.unlink(path)

    def test_skips_non_integer_question_numbers(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Pregunta", "Respuesta"])
            sheet.append([17.5, "c"])
            workbook.save(path)

            assert _parse_local_xlsx_answer_key(path) == {}
        finally:
            os.unlink(path)


class TestBuildAnswerKeyXlsx:
    def test_creates_valid_xlsx(self):
        answer_key = {1: "a", 2: "b", 3: "c"}
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            build_answer_key_xlsx(answer_key, path)
            assert os.path.exists(path)
            assert os.path.getsize(path) > 0
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            assert ws.cell(1, 1).value == "Pregunta"
            assert ws.cell(1, 2).value == "Respuesta"
            assert ws.cell(2, 1).value == 1
            assert ws.cell(2, 2).value == "a"
            assert ws.cell(3, 1).value == 2
            assert ws.cell(3, 2).value == "b"
            assert ws.cell(4, 1).value == 3
            assert ws.cell(4, 2).value == "c"
        finally:
            os.unlink(path)

    def test_sorted_by_question_number(self):
        answer_key = {3: "c", 1: "a", 2: "b"}
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            build_answer_key_xlsx(answer_key, path)
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            assert ws.cell(2, 1).value == 1
            assert ws.cell(3, 1).value == 2
            assert ws.cell(4, 1).value == 3
        finally:
            os.unlink(path)

    def test_36_questions(self):
        answer_key = {i: chr(ord("a") + (i - 1) % 4) for i in range(1, 37)}
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            build_answer_key_xlsx(answer_key, path)
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            assert ws.max_row == 37  # 1 header + 36 data
            assert ws.cell(37, 1).value == 36
            assert ws.cell(37, 2).value == "d"
        finally:
            os.unlink(path)


class TestAnswerKeyIntegration:
    def _create_exam_with_known_answers(self, path: str):
        """Crea un .docx donde la respuesta correcta siempre es 'a'."""
        from docx import Document

        doc = Document()
        doc.add_paragraph("Prueba de Física")
        doc.add_paragraph("Forma 1")
        qt = doc.add_table(rows=3, cols=2)
        for i in range(6):
            row_idx = i // 2
            col_idx = i % 2
            cell = qt.cell(row_idx, col_idx)
            cell.paragraphs[0].text = f"{i + 1}- Pregunta {i + 1}"
            cell.add_paragraph(f"a) Correcta {i + 1}")
            cell.add_paragraph(f"b) Incorrecta B")
            cell.add_paragraph(f"c) Incorrecta C")
            cell.add_paragraph(f"d) Incorrecta D")
        doc.save(path)

    def test_answer_key_tracks_correct_answer_through_shuffle(self):
        """La respuesta correcta debe seguir apuntando al texto correcto."""
        import random

        from src.parser import parse_docx
        from src.randomizer import generate_forms

        with tempfile.TemporaryDirectory() as tmpdir:
            source = os.path.join(tmpdir, "original.docx")
            self._create_exam_with_known_answers(source)

            exam = parse_docx(source)
            original_key = {i: "a" for i in range(1, 7)}

            random.seed(123)
            forms = generate_forms(exam, n=1)
            form = forms[0]

            new_key = generate_answer_key(form, original_key)

            for q in form:
                correct_letter = new_key[q.number]
                correct_alt = next(
                    a for a in q.alternatives if a.letter == correct_letter
                )
                assert "Correcta" in correct_alt.text, (
                    f"Q{q.number} (orig Q{q.original_number}): "
                    f"pauta dice '{correct_letter}' pero texto es '{correct_alt.text}'"
                )

    def test_answer_key_xlsx_roundtrip(self):
        answer_key = {1: "c", 2: "a", 3: "d", 4: "b", 5: "a", 6: "c"}
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            build_answer_key_xlsx(answer_key, path)
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            read_back = {}
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
                q_num, letter = row
                read_back[q_num] = letter
            assert read_back == answer_key
        finally:
            os.unlink(path)
