"""Tests para los 3 bug fixes:

1. Imágenes sin espacio — parser detecta párrafos con imagen como contenido (no vacío),
   doc_builder convierte anchors flotantes a inline y escala imágenes anchas.
2. Preguntas cortadas — cantSplit fuerza w:val="1" incluso si ya existe con val="0".
3. Parte final del documento cortada — contenido post-tabla se preserva tras save.
"""

import os
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree

from src.doc_builder import (
    DEFAULT_MAX_IMAGE_WIDTH_EMU,
    _build_cell_content,
    _clamp_image_width,
    _convert_anchors_to_inline,
    _fix_cell_images,
    _get_column_widths_emu,
    _set_rows_cant_split,
    build_docx,
)
from src.doc_builder import (
    _paragraph_has_image as builder_has_image,
)
from src.models import Alternative, Question
from src.parser import (
    _paragraph_has_image as parser_has_image,
)
from src.parser import (
    _parse_cell,
    parse_docx,
)
from src.randomizer import generate_forms

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_PIC = "http://schemas.openxmlformats.org/drawingml/2006/picture"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ---------------------------------------------------------------------------
# XML builders for test fixtures
# ---------------------------------------------------------------------------


def _make_empty_paragraph() -> etree._Element:
    """Crea un elemento <w:p> vacío."""
    p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
    return p


def _make_text_paragraph(text: str) -> etree._Element:
    """Crea un elemento <w:p> con texto."""
    p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
    r = etree.SubElement(p, f"{{{NS_W}}}r")
    t = etree.SubElement(r, f"{{{NS_W}}}t")
    t.text = text
    return p


def _make_anchor_image(
    cx_emu: int = 2_000_000, cy_emu: int = 1_000_000
) -> etree._Element:
    """Crea un <w:p> con un <w:drawing>/<wp:anchor> completo (imagen flotante).

    Args:
        cx_emu: Ancho de la imagen en EMUs.
        cy_emu: Alto de la imagen en EMUs.
    """
    p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
    r = etree.SubElement(p, f"{{{NS_W}}}r")
    drawing = etree.SubElement(r, f"{{{NS_W}}}drawing")

    anchor = etree.SubElement(drawing, f"{{{NS_WP}}}anchor")
    anchor.set("allowOverlap", "1")
    anchor.set("behindDoc", "0")
    anchor.set("distB", "114300")
    anchor.set("distT", "114300")
    anchor.set("distL", "114300")
    anchor.set("distR", "114300")
    anchor.set("hidden", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("locked", "0")
    anchor.set("relativeHeight", "0")
    anchor.set("simplePos", "0")

    simple_pos = etree.SubElement(anchor, f"{{{NS_WP}}}simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")

    posH = etree.SubElement(anchor, f"{{{NS_WP}}}positionH")
    posH.set("relativeFrom", "column")
    posH_off = etree.SubElement(posH, f"{{{NS_WP}}}posOffset")
    posH_off.text = "1323975"

    posV = etree.SubElement(anchor, f"{{{NS_WP}}}positionV")
    posV.set("relativeFrom", "paragraph")
    posV_off = etree.SubElement(posV, f"{{{NS_WP}}}posOffset")
    posV_off.text = "147340"

    extent = etree.SubElement(anchor, f"{{{NS_WP}}}extent")
    extent.set("cx", str(cx_emu))
    extent.set("cy", str(cy_emu))

    effect = etree.SubElement(anchor, f"{{{NS_WP}}}effectExtent")
    effect.set("b", "0")
    effect.set("l", "0")
    effect.set("r", "0")
    effect.set("t", "0")

    etree.SubElement(anchor, f"{{{NS_WP}}}wrapNone")

    docPr = etree.SubElement(anchor, f"{{{NS_WP}}}docPr")
    docPr.set("id", "1")
    docPr.set("name", "image1.png")

    graphic = etree.SubElement(anchor, f"{{{NS_A}}}graphic")
    graphicData = etree.SubElement(graphic, f"{{{NS_A}}}graphicData")
    graphicData.set("uri", NS_PIC)
    pic = etree.SubElement(graphicData, f"{{{NS_PIC}}}pic")
    nvPicPr = etree.SubElement(pic, f"{{{NS_PIC}}}nvPicPr")
    cNvPr = etree.SubElement(nvPicPr, f"{{{NS_PIC}}}cNvPr")
    cNvPr.set("id", "0")
    cNvPr.set("name", "image1.png")
    cNvPicPr = etree.SubElement(nvPicPr, f"{{{NS_PIC}}}cNvPicPr")
    cNvPicPr.set("preferRelativeResize", "0")
    blipFill = etree.SubElement(pic, f"{{{NS_PIC}}}blipFill")
    blip = etree.SubElement(blipFill, f"{{{NS_A}}}blip")
    blip.set(f"{{{NS_R}}}embed", "rId99")
    stretch = etree.SubElement(blipFill, f"{{{NS_A}}}stretch")
    etree.SubElement(stretch, f"{{{NS_A}}}fillRect")
    spPr = etree.SubElement(pic, f"{{{NS_PIC}}}spPr")
    xfrm = etree.SubElement(spPr, f"{{{NS_A}}}xfrm")
    a_off = etree.SubElement(xfrm, f"{{{NS_A}}}off")
    a_off.set("x", "0")
    a_off.set("y", "0")
    a_ext = etree.SubElement(xfrm, f"{{{NS_A}}}ext")
    a_ext.set("cx", str(cx_emu))
    a_ext.set("cy", str(cy_emu))

    return p


def _make_inline_image(
    cx_emu: int = 2_000_000, cy_emu: int = 1_000_000
) -> etree._Element:
    """Crea un <w:p> con un <w:drawing>/<wp:inline> (imagen inline)."""
    p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
    r = etree.SubElement(p, f"{{{NS_W}}}r")
    drawing = etree.SubElement(r, f"{{{NS_W}}}drawing")

    inline = etree.SubElement(drawing, f"{{{NS_WP}}}inline")
    inline.set("distB", "114300")
    inline.set("distT", "114300")
    inline.set("distL", "114300")
    inline.set("distR", "114300")

    extent = etree.SubElement(inline, f"{{{NS_WP}}}extent")
    extent.set("cx", str(cx_emu))
    extent.set("cy", str(cy_emu))

    effect = etree.SubElement(inline, f"{{{NS_WP}}}effectExtent")
    effect.set("b", "0")
    effect.set("l", "0")
    effect.set("r", "0")
    effect.set("t", "0")

    docPr = etree.SubElement(inline, f"{{{NS_WP}}}docPr")
    docPr.set("id", "1")
    docPr.set("name", "image1.png")

    graphic = etree.SubElement(inline, f"{{{NS_A}}}graphic")
    graphicData = etree.SubElement(graphic, f"{{{NS_A}}}graphicData")
    graphicData.set("uri", NS_PIC)
    pic = etree.SubElement(graphicData, f"{{{NS_PIC}}}pic")
    spPr = etree.SubElement(pic, f"{{{NS_PIC}}}spPr")
    xfrm = etree.SubElement(spPr, f"{{{NS_A}}}xfrm")
    a_ext = etree.SubElement(xfrm, f"{{{NS_A}}}ext")
    a_ext.set("cx", str(cx_emu))
    a_ext.set("cy", str(cy_emu))

    return p


def _make_paragraph_with_drawing() -> etree._Element:
    """Simple <w:p> con un <w:drawing> vacío (para tests de detección basic)."""
    p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
    r = etree.SubElement(p, f"{{{NS_W}}}r")
    etree.SubElement(r, f"{{{NS_W}}}drawing")
    return p


def _make_paragraph_with_pict() -> etree._Element:
    """Simple <w:p> con un <w:pict> legacy."""
    p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
    r = etree.SubElement(p, f"{{{NS_W}}}r")
    etree.SubElement(r, f"{{{NS_W}}}pict")
    return p


# ---------------------------------------------------------------------------
# Bug 1: Parser — detección de imágenes
# ---------------------------------------------------------------------------


class TestParserImageDetection:
    """Verifica que _paragraph_has_image detecta imágenes correctamente."""

    def test_drawing_detected(self):
        p = _make_paragraph_with_drawing()
        assert parser_has_image(p) is True

    def test_pict_detected(self):
        p = _make_paragraph_with_pict()
        assert parser_has_image(p) is True

    def test_empty_paragraph_no_image(self):
        p = _make_empty_paragraph()
        assert parser_has_image(p) is False

    def test_text_only_no_image(self):
        p = _make_text_paragraph("Solo texto, sin imagen")
        assert parser_has_image(p) is False


class TestParserImageInEnunciado:
    """Verifica que párrafos con imagen en el enunciado NO van a trailing."""

    def test_image_paragraph_stays_in_enunciado(self):
        """Un párrafo con imagen entre enunciado y alternativas debe quedar
        en enunciado_paragraphs, no en trailing_empty_paragraphs."""
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            path = f.name
        try:
            doc = Document()
            q_table = doc.add_table(rows=1, cols=1)
            cell = q_table.cell(0, 0)

            cell.paragraphs[0].text = "2- Observa la siguiente imagen:"

            # Párrafo con imagen real (inline shape)
            img_para = cell.add_paragraph()
            run = img_para.add_run()

            import io
            import struct
            import zlib

            def _make_1px_png() -> bytes:
                sig = b"\x89PNG\r\n\x1a\n"
                ihdr_data = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
                ihdr_crc = zlib.crc32(b"IHDR" + ihdr_data) & 0xFFFFFFFF
                ihdr = (
                    struct.pack(">I", 13)
                    + b"IHDR"
                    + ihdr_data
                    + struct.pack(">I", ihdr_crc)
                )
                raw = b"\x00\x00\x00\x00"
                idat_data = zlib.compress(raw)
                idat_crc = zlib.crc32(b"IDAT" + idat_data) & 0xFFFFFFFF
                idat = (
                    struct.pack(">I", len(idat_data))
                    + b"IDAT"
                    + idat_data
                    + struct.pack(">I", idat_crc)
                )
                iend_crc = zlib.crc32(b"IEND") & 0xFFFFFFFF
                iend = struct.pack(">I", 0) + b"IEND" + struct.pack(">I", iend_crc)
                return sig + ihdr + idat + iend

            png_bytes = _make_1px_png()
            run.add_picture(io.BytesIO(png_bytes), width=Inches(2))

            cell.add_paragraph("a) Transversal")
            cell.add_paragraph("b) Longitudinal")
            cell.add_paragraph("c) Electromagnética")
            cell.add_paragraph("d) Mecánica")

            doc.save(path)

            exam = parse_docx(path)
            assert len(exam.questions) == 1

            q = exam.questions[0]
            assert q.number == 2
            assert len(q.enunciado_paragraphs) == 2, (
                f"Esperaba 2 párrafos de enunciado (texto + imagen), "
                f"got {len(q.enunciado_paragraphs)}"
            )
            assert len(q.trailing_empty_paragraphs) == 0, (
                f"La imagen no debería ir a trailing, "
                f"pero hay {len(q.trailing_empty_paragraphs)} trailing paragraphs"
            )
            assert len(q.alternatives) == 4
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# Bug 1: Builder — detección de imágenes
# ---------------------------------------------------------------------------


class TestBuilderImageDetection:
    """Verifica que doc_builder detecta imágenes en párrafos."""

    def test_drawing_detected(self):
        p = _make_paragraph_with_drawing()
        assert builder_has_image(p) is True

    def test_empty_not_detected(self):
        p = _make_empty_paragraph()
        assert builder_has_image(p) is False


# ---------------------------------------------------------------------------
# Bug 1: Conversión anchor → inline
# ---------------------------------------------------------------------------


class TestConvertAnchorsToInline:
    """Verifica que _convert_anchors_to_inline convierte correctamente."""

    def test_anchor_becomes_inline(self):
        p = _make_anchor_image(cx_emu=1_668_145, cy_emu=971_550)
        # Antes: hay un anchor
        anchors = p.findall(f".//{{{NS_WP}}}anchor")
        assert len(anchors) == 1

        _convert_anchors_to_inline(p)

        # Después: no hay anchors, hay un inline
        anchors_after = p.findall(f".//{{{NS_WP}}}anchor")
        assert len(anchors_after) == 0
        inlines_after = p.findall(f".//{{{NS_WP}}}inline")
        assert len(inlines_after) == 1

    def test_removes_anchor_only_attributes(self):
        p = _make_anchor_image()
        _convert_anchors_to_inline(p)

        inline = p.findall(f".//{{{NS_WP}}}inline")[0]
        for attr in (
            "allowOverlap",
            "behindDoc",
            "hidden",
            "layoutInCell",
            "locked",
            "relativeHeight",
            "simplePos",
        ):
            assert attr not in inline.attrib, (
                f"Atributo '{attr}' debería haberse removido"
            )

    def test_removes_position_elements(self):
        p = _make_anchor_image()
        _convert_anchors_to_inline(p)

        inline = p.findall(f".//{{{NS_WP}}}inline")[0]
        assert inline.find(f"{{{NS_WP}}}positionH") is None
        assert inline.find(f"{{{NS_WP}}}positionV") is None
        assert inline.find(f"{{{NS_WP}}}wrapNone") is None
        assert inline.find(f"{{{NS_WP}}}simplePos") is None

    def test_preserves_extent(self):
        cx, cy = 1_668_145, 971_550
        p = _make_anchor_image(cx_emu=cx, cy_emu=cy)
        _convert_anchors_to_inline(p)

        inline = p.findall(f".//{{{NS_WP}}}inline")[0]
        extent = inline.find(f"{{{NS_WP}}}extent")
        assert extent is not None
        assert extent.get("cx") == str(cx)
        assert extent.get("cy") == str(cy)

    def test_preserves_graphic(self):
        p = _make_anchor_image()
        _convert_anchors_to_inline(p)

        inline = p.findall(f".//{{{NS_WP}}}inline")[0]
        graphic = inline.find(f"{{{NS_A}}}graphic")
        assert graphic is not None

    def test_preserves_dist_attributes(self):
        p = _make_anchor_image()
        _convert_anchors_to_inline(p)

        inline = p.findall(f".//{{{NS_WP}}}inline")[0]
        assert inline.get("distB") == "114300"
        assert inline.get("distT") == "114300"

    def test_inline_image_unaffected(self):
        """Una imagen que ya es inline no debe cambiar."""
        p = _make_inline_image(cx_emu=2_000_000, cy_emu=1_000_000)
        inlines_before = p.findall(f".//{{{NS_WP}}}inline")
        assert len(inlines_before) == 1

        _convert_anchors_to_inline(p)

        inlines_after = p.findall(f".//{{{NS_WP}}}inline")
        assert len(inlines_after) == 1
        # Extent sin cambios
        extent = inlines_after[0].find(f"{{{NS_WP}}}extent")
        assert extent.get("cx") == "2000000"

    def test_no_images_no_crash(self):
        p = _make_text_paragraph("Solo texto")
        _convert_anchors_to_inline(p)  # No debería fallar

    def test_multiple_anchors_converted(self):
        """Si un párrafo tiene múltiples anchors, todos se convierten."""
        # Construir párrafo con 2 anchors
        p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
        for i in range(2):
            r = etree.SubElement(p, f"{{{NS_W}}}r")
            drawing = etree.SubElement(r, f"{{{NS_W}}}drawing")
            anchor = etree.SubElement(drawing, f"{{{NS_WP}}}anchor")
            anchor.set("simplePos", "0")
            anchor.set("allowOverlap", "1")
            etree.SubElement(anchor, f"{{{NS_WP}}}simplePos").set("x", "0")
            etree.SubElement(anchor, f"{{{NS_WP}}}wrapNone")
            extent = etree.SubElement(anchor, f"{{{NS_WP}}}extent")
            extent.set("cx", str(1_000_000 * (i + 1)))
            extent.set("cy", "500000")

        assert len(p.findall(f".//{{{NS_WP}}}anchor")) == 2

        _convert_anchors_to_inline(p)

        assert len(p.findall(f".//{{{NS_WP}}}anchor")) == 0
        assert len(p.findall(f".//{{{NS_WP}}}inline")) == 2


# ---------------------------------------------------------------------------
# Bug 1: Clamp de ancho de imagen
# ---------------------------------------------------------------------------


class TestClampImageWidth:
    """Verifica que _clamp_image_width escala imágenes correctamente."""

    def test_oversized_image_clamped(self):
        """Una imagen más ancha que el máximo se escala proporcionalmente."""
        cx = 4_000_000  # Muy ancha
        cy = 2_000_000
        max_w = 3_000_000
        p = _make_inline_image(cx_emu=cx, cy_emu=cy)

        _clamp_image_width(p, max_w)

        extent = p.findall(f".//{{{NS_WP}}}extent")[0]
        assert int(extent.get("cx")) == max_w
        # Aspect ratio preservado: cy_new = 2_000_000 * (3_000_000 / 4_000_000) = 1_500_000
        assert int(extent.get("cy")) == 1_500_000

    def test_undersized_image_untouched(self):
        """Una imagen más angosta que el máximo no cambia."""
        cx = 1_500_000
        cy = 800_000
        max_w = 3_000_000
        p = _make_inline_image(cx_emu=cx, cy_emu=cy)

        _clamp_image_width(p, max_w)

        extent = p.findall(f".//{{{NS_WP}}}extent")[0]
        assert int(extent.get("cx")) == cx
        assert int(extent.get("cy")) == cy

    def test_exact_width_untouched(self):
        """Una imagen con ancho exactamente igual al máximo no cambia."""
        cx = 3_000_000
        cy = 1_000_000
        p = _make_inline_image(cx_emu=cx, cy_emu=cy)

        _clamp_image_width(p, cx)

        extent = p.findall(f".//{{{NS_WP}}}extent")[0]
        assert int(extent.get("cx")) == cx
        assert int(extent.get("cy")) == cy

    def test_a_ext_also_scaled(self):
        """El <a:ext> interno de la imagen también se escala."""
        cx = 4_000_000
        cy = 2_000_000
        max_w = 2_000_000
        p = _make_anchor_image(cx_emu=cx, cy_emu=cy)
        # Convertir primero para que sea inline
        _convert_anchors_to_inline(p)

        _clamp_image_width(p, max_w)

        a_exts = p.findall(f".//{{{NS_A}}}ext")
        assert len(a_exts) > 0
        a_ext = a_exts[0]
        # scale = 2_000_000 / 4_000_000 = 0.5
        assert int(a_ext.get("cx")) == 2_000_000
        assert int(a_ext.get("cy")) == 1_000_000

    def test_no_images_no_crash(self):
        p = _make_text_paragraph("Solo texto")
        _clamp_image_width(p, 3_000_000)  # No debería fallar


# ---------------------------------------------------------------------------
# Bug 1: _fix_cell_images integración
# ---------------------------------------------------------------------------


class TestFixCellImages:
    """Verifica que _fix_cell_images convierte y escala en un solo paso."""

    def test_anchor_converted_and_clamped(self):
        """Un anchor demasiado ancho se convierte a inline y se escala."""
        max_w = 2_500_000
        p = _make_anchor_image(cx_emu=3_500_000, cy_emu=1_500_000)

        _fix_cell_images([p], max_w)

        # Convertido a inline
        assert len(p.findall(f".//{{{NS_WP}}}anchor")) == 0
        assert len(p.findall(f".//{{{NS_WP}}}inline")) == 1
        # Escalado
        extent = p.findall(f".//{{{NS_WP}}}extent")[0]
        assert int(extent.get("cx")) == max_w

    def test_mixed_paragraphs(self):
        """Mezcla de texto y anchors: solo los anchors se procesan."""
        text_p = _make_text_paragraph("1- Pregunta")
        img_p = _make_anchor_image(cx_emu=2_000_000, cy_emu=1_000_000)
        alt_p = _make_text_paragraph("a) Alternativa")

        _fix_cell_images([text_p, img_p, alt_p], 3_000_000)

        # Solo img_p cambió
        assert len(img_p.findall(f".//{{{NS_WP}}}inline")) == 1
        assert len(img_p.findall(f".//{{{NS_WP}}}anchor")) == 0
        # text y alt no cambiaron
        runs = text_p.findall(f".//{{{NS_W}}}t")
        assert runs[0].text == "1- Pregunta"


# ---------------------------------------------------------------------------
# Bug 1: _build_cell_content integración
# ---------------------------------------------------------------------------


class TestBuildCellContentWithImages:
    """Verifica que _build_cell_content convierte y escala imágenes."""

    def test_anchor_in_enunciado_converted_to_inline(self):
        """Un anchor en el enunciado se convierte a inline en el output."""
        text_para = _make_text_paragraph("1- Observa:")
        img_para = _make_anchor_image(cx_emu=2_000_000, cy_emu=1_000_000)

        q = Question(
            number=1,
            enunciado_paragraphs=[text_para, img_para],
            alternatives=[
                Alternative(letter="a", paragraphs=[_make_text_paragraph("a) Sí")]),
                Alternative(letter="b", paragraphs=[_make_text_paragraph("b) No")]),
            ],
            trailing_empty_paragraphs=[],
        )

        paragraphs = _build_cell_content(q)
        img_p = paragraphs[1]
        assert len(img_p.findall(f".//{{{NS_WP}}}inline")) == 1
        assert len(img_p.findall(f".//{{{NS_WP}}}anchor")) == 0

    def test_oversized_image_clamped_in_build(self):
        """Una imagen más ancha que max_image_width_emu se escala."""
        max_w = 2_000_000
        img_para = _make_anchor_image(cx_emu=3_500_000, cy_emu=1_500_000)

        q = Question(
            number=1,
            enunciado_paragraphs=[_make_text_paragraph("1- Mira:"), img_para],
            alternatives=[
                Alternative(letter="a", paragraphs=[_make_text_paragraph("a) X")]),
            ],
            trailing_empty_paragraphs=[],
        )

        paragraphs = _build_cell_content(q, max_image_width_emu=max_w)
        img_p = paragraphs[1]
        extent = img_p.findall(f".//{{{NS_WP}}}extent")[0]
        assert int(extent.get("cx")) == max_w


# ---------------------------------------------------------------------------
# Bug 1: _get_column_widths_emu
# ---------------------------------------------------------------------------


class TestGetColumnWidthsEmu:
    """Verifica extracción de anchos de columna desde tblGrid."""

    def test_reads_grid_cols(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        # python-docx crea tblGrid automáticamente
        widths = _get_column_widths_emu(table)
        # Debe retornar 2 anchos (uno por columna)
        assert len(widths) == 2
        for w in widths:
            assert w > 0

    def test_empty_table_returns_values(self):
        doc = Document()
        table = doc.add_table(rows=0, cols=1)
        widths = _get_column_widths_emu(table)
        # Puede retornar widths o vacío dependiendo de python-docx
        assert isinstance(widths, list)


# ---------------------------------------------------------------------------
# Bug 2: cantSplit forzando val="1"
# ---------------------------------------------------------------------------


class TestCantSplitOverride:
    """Verifica que _set_rows_cant_split fuerza w:val='1'."""

    def test_new_cant_split_has_val_1(self):
        """En fila sin cantSplit, lo crea con val='1'."""
        doc = Document()
        table = doc.add_table(rows=2, cols=1)

        _set_rows_cant_split(table)

        for row in table.rows:
            tr = row._tr
            trPr = tr.find(qn("w:trPr"))
            assert trPr is not None
            cant_split = trPr.find(qn("w:cantSplit"))
            assert cant_split is not None
            assert cant_split.get(qn("w:val")) == "1"

    def test_overrides_existing_val_0(self):
        """Si cantSplit existe con val='0', lo fuerza a '1'."""
        doc = Document()
        table = doc.add_table(rows=1, cols=1)

        tr = table.rows[0]._tr
        trPr = OxmlElement("w:trPr")
        tr.insert(0, trPr)
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "0")
        trPr.append(cant_split)

        _set_rows_cant_split(table)

        cant_split_after = trPr.find(qn("w:cantSplit"))
        assert cant_split_after.get(qn("w:val")) == "1", (
            "cantSplit debería ser '1' después de override"
        )

    def test_preserves_existing_val_1(self):
        """Si cantSplit ya tiene val='1', no lo rompe."""
        doc = Document()
        table = doc.add_table(rows=1, cols=1)

        tr = table.rows[0]._tr
        trPr = OxmlElement("w:trPr")
        tr.insert(0, trPr)
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "1")
        trPr.append(cant_split)

        _set_rows_cant_split(table)

        cant_split_after = trPr.find(qn("w:cantSplit"))
        assert cant_split_after.get(qn("w:val")) == "1"


# ---------------------------------------------------------------------------
# Bug 3: Preservación de contenido post-tabla
# ---------------------------------------------------------------------------


class TestPostTableContentPreservation:
    """Verifica que el contenido después de la tabla de preguntas se preserva."""

    def _create_exam_with_footer(self, path: str):
        """Crea un .docx con contenido después de la tabla de preguntas."""
        doc = Document()

        p = doc.add_paragraph()
        run = p.add_run("Forma 1")
        run.bold = True

        ht = doc.add_table(rows=1, cols=2)
        ht.cell(0, 0).text = "NOMBRE"
        ht.cell(0, 1).text = ""

        it = doc.add_table(rows=1, cols=1)
        it.cell(0, 0).text = "INSTRUCCIONES"

        qt = doc.add_table(rows=2, cols=2)
        for i in range(4):
            row_idx = i // 2
            col_idx = i % 2
            cell = qt.cell(row_idx, col_idx)
            cell.paragraphs[0].text = f"{i + 1}- Pregunta {i + 1}"
            cell.add_paragraph("a) Opción A")
            cell.add_paragraph("b) Opción B")
            cell.add_paragraph("c) Opción C")
            cell.add_paragraph("d) Opción D")

        footer_p = doc.add_paragraph()
        footer_run = footer_p.add_run("¡Has llegado al final de la prueba!")
        footer_run.bold = True

        footer_p2 = doc.add_paragraph()
        footer_p2.add_run(
            "¡Recuerda pasar tus respuestas a la hoja de respuestas "
            "hasta la pregunta 36!"
        )

        doc.save(path)

    def test_footer_text_preserved_after_build(self):
        """El texto después de la tabla de preguntas debe permanecer en el output."""
        with tempfile.TemporaryDirectory() as tmpdir:
            source = os.path.join(tmpdir, "original.docx")
            output = os.path.join(tmpdir, "forma_2.docx")
            self._create_exam_with_footer(source)

            exam = parse_docx(source)
            forms = generate_forms(exam, n=1)

            build_docx(
                source_path=source,
                questions=forms[0],
                output_path=output,
                form_number=2,
                questions_table_index=exam.questions_table_index,
                questions_per_row=exam.questions_per_row,
            )

            out_doc = Document(output)
            all_text = " ".join(p.text for p in out_doc.paragraphs)

            assert "final de la prueba" in all_text, (
                f"El texto de cierre se perdió. Texto completo de párrafos: {all_text}"
            )
            assert "hoja de respuestas" in all_text, (
                f"Las instrucciones finales se perdieron. Texto completo: {all_text}"
            )

    def test_footer_preserved_with_multiple_forms(self):
        """El contenido post-tabla se preserva en todas las formas generadas."""
        with tempfile.TemporaryDirectory() as tmpdir:
            source = os.path.join(tmpdir, "original.docx")
            self._create_exam_with_footer(source)

            exam = parse_docx(source)
            forms = generate_forms(exam, n=3)

            for i, form_questions in enumerate(forms):
                output = os.path.join(tmpdir, f"forma_{i + 2}.docx")
                build_docx(
                    source_path=source,
                    questions=form_questions,
                    output_path=output,
                    form_number=i + 2,
                    questions_table_index=exam.questions_table_index,
                    questions_per_row=exam.questions_per_row,
                )

                out_doc = Document(output)
                all_text = " ".join(p.text for p in out_doc.paragraphs)
                assert "final de la prueba" in all_text, (
                    f"Forma {i + 2}: texto de cierre perdido"
                )

    def test_source_footer_paragraph_count(self):
        """Verifica que el documento fuente tiene párrafos después de la tabla."""
        with tempfile.TemporaryDirectory() as tmpdir:
            source = os.path.join(tmpdir, "original.docx")
            self._create_exam_with_footer(source)

            doc = Document(source)
            footer_texts = [
                p.text
                for p in doc.paragraphs
                if "final" in p.text.lower() or "respuestas" in p.text.lower()
            ]
            assert len(footer_texts) == 2, (
                f"El documento fuente debe tener 2 párrafos de footer, "
                f"encontrados: {footer_texts}"
            )


# ---------------------------------------------------------------------------
# Bug 1 (follow-up): Separación de párrafos mixtos texto + imagen
# ---------------------------------------------------------------------------


from src.doc_builder import _split_text_and_image_runs


def _make_mixed_paragraph(text: str, cx_emu: int = 1_000_000) -> etree._Element:
    """Crea un <w:p> con un run de texto Y un run con <w:drawing> (caso mixto real).

    Estructura:
      <w:p>
        <w:pPr/>
        <w:r><w:t>texto</w:t></w:r>
        <w:r><w:drawing>...</w:drawing></w:r>
    """
    p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
    # pPr
    etree.SubElement(p, f"{{{NS_W}}}pPr")
    # run de texto
    r_text = etree.SubElement(p, f"{{{NS_W}}}r")
    t = etree.SubElement(r_text, f"{{{NS_W}}}t")
    t.text = text
    # run de imagen con drawing
    r_img = etree.SubElement(p, f"{{{NS_W}}}r")
    drawing = etree.SubElement(r_img, f"{{{NS_W}}}drawing")
    inline = etree.SubElement(drawing, f"{{{NS_WP}}}inline")
    inline.set("distB", "114300")
    extent = etree.SubElement(inline, f"{{{NS_WP}}}extent")
    extent.set("cx", str(cx_emu))
    extent.set("cy", "500000")
    return p


class TestSplitTextAndImageRuns:
    """Verifica que _split_text_and_image_runs separa párrafos mixtos correctamente."""

    def test_pure_text_paragraph_unchanged(self):
        """Un párrafo solo con texto se retorna tal cual (misma referencia)."""
        p = _make_text_paragraph("1- Enunciado de la pregunta")
        result = _split_text_and_image_runs(p)
        assert len(result) == 1
        assert result[0] is p  # misma referencia, sin copiar

    def test_pure_image_paragraph_unchanged(self):
        """Un párrafo solo con imagen se retorna tal cual (misma referencia)."""
        p = _make_inline_image(cx_emu=1_500_000, cy_emu=800_000)
        result = _split_text_and_image_runs(p)
        assert len(result) == 1
        assert result[0] is p

    def test_empty_paragraph_unchanged(self):
        """Un párrafo vacío se retorna tal cual."""
        p = _make_empty_paragraph()
        result = _split_text_and_image_runs(p)
        assert len(result) == 1
        assert result[0] is p

    def test_mixed_paragraph_splits_into_two(self):
        """Un párrafo con texto e imagen se divide en 2 párrafos."""
        p = _make_mixed_paragraph("22- ¿Cuál de las ondas tiene mayor amplitud?")
        result = _split_text_and_image_runs(p)
        assert len(result) == 2

    def test_first_result_has_text_no_image(self):
        """El primer párrafo resultante contiene el texto pero no la imagen."""
        p = _make_mixed_paragraph("22- Texto del enunciado")
        text_p, img_p = _split_text_and_image_runs(p)

        # Tiene texto
        texts = text_p.findall(f".//{{{NS_W}}}t")
        assert any(t.text and "Texto del enunciado" in t.text for t in texts)

        # No tiene imagen
        assert len(text_p.findall(f".//{{{NS_W}}}drawing")) == 0

    def test_second_result_has_image_no_text(self):
        """El segundo párrafo resultante contiene la imagen pero no el texto."""
        p = _make_mixed_paragraph("22- Texto del enunciado")
        text_p, img_p = _split_text_and_image_runs(p)

        # Tiene imagen
        assert len(img_p.findall(f".//{{{NS_W}}}drawing")) == 1

        # No tiene texto visible
        texts = img_p.findall(f".//{{{NS_W}}}t")
        assert not any(t.text and t.text.strip() for t in texts)

    def test_pPr_copied_to_both_paragraphs(self):
        """<w:pPr> se copia tanto al párrafo de texto como al de imagen."""
        p = _make_mixed_paragraph("23- Enunciado con imagen", cx_emu=2_000_000)
        text_p, img_p = _split_text_and_image_runs(p)

        assert text_p.find(f"{{{NS_W}}}pPr") is not None, "text_p debe tener <w:pPr>"
        assert img_p.find(f"{{{NS_W}}}pPr") is not None, "img_p debe tener <w:pPr>"

    def test_original_paragraph_not_mutated(self):
        """El párrafo original no debe modificarse (los nuevos son copias)."""
        p = _make_mixed_paragraph("Original intacto")
        children_before = len(list(p))
        _split_text_and_image_runs(p)
        children_after = len(list(p))
        assert children_before == children_after, (
            "Los hijos del párrafo original no deben cambiar"
        )

    def test_multiple_text_runs_all_go_to_text_paragraph(self):
        """Si hay múltiples runs de texto, todos van al párrafo de texto."""
        p = etree.SubElement(etree.Element("root"), f"{{{NS_W}}}p")
        # 2 runs de texto
        for txt in ("25- Parte 1 ", "del enunciado"):
            r = etree.SubElement(p, f"{{{NS_W}}}r")
            t = etree.SubElement(r, f"{{{NS_W}}}t")
            t.text = txt
        # 1 run de imagen
        r_img = etree.SubElement(p, f"{{{NS_W}}}r")
        drawing = etree.SubElement(r_img, f"{{{NS_W}}}drawing")
        etree.SubElement(drawing, f"{{{NS_WP}}}inline")

        text_p, img_p = _split_text_and_image_runs(p)

        text_runs = text_p.findall(f"{{{NS_W}}}r")
        assert len(text_runs) == 2
        img_runs = img_p.findall(f"{{{NS_W}}}r")
        assert len(img_runs) == 1

    def test_integration_with_build_cell_content(self):
        """_build_cell_content separa párrafos mixtos y pone la imagen al final del enunciado."""
        from src.doc_builder import _build_cell_content

        mixed_p = _make_mixed_paragraph("1- Enunciado con imagen adjunta")

        q = Question(
            number=1,
            enunciado_paragraphs=[mixed_p],
            alternatives=[
                Alternative(letter="a", paragraphs=[_make_text_paragraph("a) Sí")]),
                Alternative(letter="b", paragraphs=[_make_text_paragraph("b) No")]),
            ],
            trailing_empty_paragraphs=[],
        )

        paragraphs = _build_cell_content(q)

        # El párrafo mixto debe haberse separado en 2
        # Total: 2 (enunciado separado) + 2 (alternativas) = 4
        assert len(paragraphs) == 4, (
            f"Esperaba 4 párrafos (2 enunciado + 2 alts), got {len(paragraphs)}"
        )

        # Texto primero, imagen después (siempre, independiente de posición original)
        first_has_text = any(
            t.text and "Enunciado" in t.text
            for t in paragraphs[0].findall(f".//{{{NS_W}}}t")
        )
        second_has_image = bool(paragraphs[1].findall(f".//{{{NS_W}}}drawing"))
        assert first_has_text, "Primer párrafo debe ser el texto del enunciado"
        assert second_has_image, (
            "Segundo párrafo debe ser la imagen (siempre debajo del texto)"
        )

    def test_image_at_start_of_enunciado_moved_to_end(self):
        """Una imagen que estaba flotando al inicio/medio del enunciado queda al final."""
        from src.doc_builder import _build_cell_content

        # Imagen primero, luego texto (caso donde el anchor flotaba arriba)
        img_p = _make_inline_image(cx_emu=1_500_000, cy_emu=800_000)
        text_p = _make_text_paragraph("5- Pregunta con imagen al inicio original")

        q = Question(
            number=5,
            enunciado_paragraphs=[img_p, text_p],  # imagen primero intencionalmente
            alternatives=[
                Alternative(letter="a", paragraphs=[_make_text_paragraph("a) A")]),
                Alternative(letter="b", paragraphs=[_make_text_paragraph("b) B")]),
            ],
            trailing_empty_paragraphs=[],
        )

        paragraphs = _build_cell_content(q)

        # Total: 1 texto + 1 imagen (enunciado) + 2 alternativas = 4
        assert len(paragraphs) == 4

        # El primer párrafo del enunciado debe ser el texto, no la imagen
        first_has_text = any(
            t.text and "Pregunta con imagen" in t.text
            for t in paragraphs[0].findall(f".//{{{NS_W}}}t")
        )
        second_has_image = bool(
            paragraphs[1].findall(f".//{{{NS_WP}}}inline")
            or paragraphs[1].findall(f".//{{{NS_W}}}drawing")
        )
        assert first_has_text, "El texto del enunciado debe ir ANTES de la imagen"
        assert second_has_image, "La imagen debe ir DESPUÉS del texto del enunciado"

    def test_multiple_text_paragraphs_then_image(self):
        """Con enunciado multilinea + imagen flotante, todo el texto va antes de la imagen."""
        from src.doc_builder import _build_cell_content

        text1 = _make_text_paragraph("13- Observa el siguiente gráfico")
        img_p = _make_inline_image(cx_emu=2_000_000, cy_emu=1_000_000)
        text2 = _make_text_paragraph("que muestra una onda transversal:")

        q = Question(
            number=13,
            # imagen intercalada entre los párrafos de texto
            enunciado_paragraphs=[text1, img_p, text2],
            alternatives=[
                Alternative(letter="a", paragraphs=[_make_text_paragraph("a) A")]),
            ],
            trailing_empty_paragraphs=[],
        )

        paragraphs = _build_cell_content(q)

        # Total: 2 texto + 1 imagen + 1 alternativa = 4
        assert len(paragraphs) == 4

        # Verificar que ningún párrafo de texto viene después de la imagen
        # dentro del bloque del enunciado (primeros 3 párrafos)
        enunciado_block = paragraphs[:3]
        found_image = False
        for p in enunciado_block:
            is_image = bool(
                p.findall(f".//{{{NS_WP}}}inline") or p.findall(f".//{{{NS_W}}}drawing")
            )
            has_visible_text = any(
                t.text and t.text.strip() for t in p.findall(f".//{{{NS_W}}}t")
            )
            if found_image:
                assert not has_visible_text, (
                    "Ningún párrafo de texto debe venir después de la imagen en el enunciado"
                )
            if is_image:
                found_image = True

        assert found_image, (
            "Debe haber exactamente una imagen en el bloque del enunciado"
        )
