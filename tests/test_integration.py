"""Test de integración end-to-end — verifica el flujo completo sin Drive.

Crea un .docx con estructura de tabla realista, lo parsea, genera formas,
y verifica que los documentos generados preservan formato y contenido.
"""

import os
import tempfile

from docx import Document
from docx.shared import Inches, Pt

from src.doc_builder import build_docx
from src.parser import parse_docx
from src.randomizer import generate_forms


def _create_realistic_exam(path: str):
    """Crea un documento .docx con estructura realista de examen con tabla."""
    doc = Document()

    # Título con formato
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("Prueba 1. Ondas y sus características")
    run_title.bold = True
    run_title.font.size = Pt(14)

    # Forma
    p_form = doc.add_paragraph()
    run_form = p_form.add_run("Forma 1")
    run_form.bold = True
    run_form.font.size = Pt(13)

    # Tabla datos alumno
    data_table = doc.add_table(rows=2, cols=4)
    data_table.cell(0, 0).paragraphs[0].add_run("NOMBRE").bold = True
    data_table.cell(0, 2).paragraphs[0].add_run("CURSO").bold = True
    data_table.cell(1, 0).paragraphs[0].add_run("FECHA").bold = True
    data_table.cell(1, 2).paragraphs[0].add_run("PUNTAJE").bold = True

    # Tabla instrucciones
    inst_table = doc.add_table(rows=1, cols=1)
    inst_cell = inst_table.cell(0, 0)
    inst_cell.paragraphs[0].add_run("INSTRUCCIONES:").bold = True
    inst_cell.add_paragraph("Lea con atención y seleccione su alternativa correcta.")

    # Párrafo de ítem
    p_item = doc.add_paragraph()
    r1 = p_item.add_run("Ítem I")
    r1.bold = True
    r1.italic = True
    r2 = p_item.add_run(". Marque la alternativa según corresponda (1 pto c/u)")
    r2.bold = True

    # Tabla de preguntas (5 filas x 2 cols = 10 preguntas)
    qt = doc.add_table(rows=5, cols=2)

    questions = [
        (
            "¿Cuál es la velocidad de la luz?",
            ["3x10^8 m/s", "3x10^6 m/s", "3x10^10 m/s", "3x10^4 m/s"],
        ),
        ("¿Qué transmiten las ondas?", ["Materia", "Fuerza", "Energía", "Calor"]),
        (
            "¿Qué es el punto de equilibrio?",
            ["Posición de reposo", "Amplitud", "Valle", "Cresta"],
        ),
        (
            "¿Qué representa la amplitud?",
            ["Distancia al equilibrio", "2 valles", "2 montes", "Tiempo"],
        ),
        (
            "¿Qué es un monte?",
            ["Punto más alto", "Punto más bajo", "Equilibrio", "Nodo"],
        ),
        ("¿Qué es un valle?", ["Punto más bajo", "Punto más alto", "Amplitud", "Fase"]),
        (
            "¿Qué tipo de onda es el sonido?",
            ["Longitudinal", "Transversal", "Electromagnética", "Estacionaria"],
        ),
        (
            "¿Qué necesitan las ondas mecánicas?",
            ["Medio material", "Vacío", "Luz", "Electricidad"],
        ),
        (
            "¿Qué es un pulso?",
            ["Perturbación aislada", "Onda periódica", "Vibración", "Resonancia"],
        ),
        (
            "¿Cómo se propaga la luz en el vacío?",
            ["Como onda EM", "No se propaga", "Como sonido", "Como calor"],
        ),
    ]

    for i, (q_text, alts) in enumerate(questions):
        row_idx = i // 2
        col_idx = i % 2
        cell = qt.cell(row_idx, col_idx)
        # Enunciado con alguna palabra en negrita
        p = cell.paragraphs[0]
        p.add_run(f"{i + 1}- ")
        if i % 3 == 0:
            # Alguna pregunta con negrita parcial
            words = q_text.split(" ")
            mid = len(words) // 2
            p.add_run(" ".join(words[:mid]))
            bold_run = p.add_run(f" {words[mid]} ")
            bold_run.bold = True
            p.add_run(" ".join(words[mid + 1 :]))
        else:
            p.add_run(q_text)

        for j, alt_text in enumerate(alts):
            letter = chr(ord("a") + j)
            cell.add_paragraph(f"{letter}) {alt_text}")

    # Footer
    doc.add_paragraph("")
    p_footer = doc.add_paragraph()
    r_footer = p_footer.add_run("¡Has llegado al final de la prueba!")
    r_footer.bold = True
    r_footer.font.size = Pt(15)

    doc.save(path)


def test_full_pipeline():
    """Test completo: crear docx → parsear → randomizar → generar formas."""
    with tempfile.TemporaryDirectory() as tmpdir:
        source_path = os.path.join(tmpdir, "original.docx")
        _create_realistic_exam(source_path)

        # Parsear
        exam = parse_docx(source_path)
        assert len(exam.questions) == 10
        assert exam.questions_per_row == 2

        # Generar 3 formas
        forms = generate_forms(exam, n=3)
        assert len(forms) == 3

        # Construir archivos
        output_paths = []
        for i, form_questions in enumerate(forms):
            form_num = 2 + i
            out_path = os.path.join(tmpdir, f"original F{form_num}.docx")
            build_docx(
                source_path=source_path,
                questions=form_questions,
                output_path=out_path,
                form_number=form_num,
                questions_table_index=exam.questions_table_index,
                questions_per_row=exam.questions_per_row,
            )
            output_paths.append(out_path)

        # Verificar que los archivos existen
        for out_path in output_paths:
            assert os.path.exists(out_path)
            assert os.path.getsize(out_path) > 0

        # Re-parsear para verificar contenido
        for out_path in output_paths:
            out_exam = parse_docx(out_path)
            assert len(out_exam.questions) == 10
            for q in out_exam.questions:
                assert len(q.alternatives) == 4

        # Verificar que las formas tienen orden distinto
        form_orders = []
        for out_path in output_paths:
            parsed = parse_docx(out_path)
            order = tuple(q.text for q in parsed.questions)
            form_orders.append(order)
        unique_orders = set(form_orders)
        assert len(unique_orders) >= 2


def test_preserves_header_and_footer():
    """Verifica que el header, tablas de datos e instrucciones se preservan."""
    with tempfile.TemporaryDirectory() as tmpdir:
        source_path = os.path.join(tmpdir, "original.docx")
        _create_realistic_exam(source_path)

        exam = parse_docx(source_path)
        forms = generate_forms(exam, n=1)

        out_path = os.path.join(tmpdir, "forma_f2.docx")
        build_docx(
            source_path=source_path,
            questions=forms[0],
            output_path=out_path,
            form_number=2,
            questions_table_index=exam.questions_table_index,
            questions_per_row=exam.questions_per_row,
        )

        out_doc = Document(out_path)

        # Header: título debe existir
        all_text = " ".join(p.text for p in out_doc.paragraphs)
        assert "Prueba 1" in all_text
        assert "Forma 2" in all_text
        assert "final de la prueba" in all_text

        # Tabla de datos del alumno
        assert "NOMBRE" in out_doc.tables[0].cell(0, 0).text
        assert "CURSO" in out_doc.tables[0].cell(0, 2).text

        # Tabla de instrucciones
        assert "INSTRUCCIONES" in out_doc.tables[1].cell(0, 0).text


def test_preserves_formatting_in_output():
    """Verifica que negrita/cursiva se preservan en el output."""
    with tempfile.TemporaryDirectory() as tmpdir:
        source_path = os.path.join(tmpdir, "original.docx")
        _create_realistic_exam(source_path)

        exam = parse_docx(source_path)
        # Clonar sin randomizar para poder verificar formato
        questions = [q.clone() for q in exam.questions]

        out_path = os.path.join(tmpdir, "forma_f2.docx")
        build_docx(
            source_path=source_path,
            questions=questions,
            output_path=out_path,
            form_number=2,
            questions_table_index=exam.questions_table_index,
            questions_per_row=exam.questions_per_row,
        )

        out_doc = Document(out_path)

        # El título debe tener bold
        title_para = out_doc.paragraphs[0]
        assert any(run.bold for run in title_para.runs)

        # Verificar que alguna pregunta en la tabla tiene bold parcial
        qt = out_doc.tables[exam.questions_table_index]
        first_cell = qt.cell(0, 0)
        first_para = first_cell.paragraphs[0]
        has_bold_run = any(run.bold for run in first_para.runs if run.text.strip())
        has_non_bold_run = any(
            not run.bold for run in first_para.runs if run.text.strip()
        )
        # Debe tener mezcla de bold y no-bold (formato parcial)
        assert has_bold_run and has_non_bold_run
