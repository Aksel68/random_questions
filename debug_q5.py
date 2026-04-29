#!/usr/bin/env python3
"""Script de diagnóstico para inspeccionar la estructura de la pregunta 5.

Uso:
    # Con archivo local:
    python debug_q5.py --local /ruta/al/archivo.docx

    # Descargando desde Drive:
    python debug_q5.py --drive "Carpeta/Ondas 2026. Ev. 1 F2.docx"

    # Número de pregunta personalizado:
    python debug_q5.py --local /ruta/archivo.docx --question 5
"""

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from lxml import etree

NSMAP = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"

import re

QUESTION_PATTERN = re.compile(r"^(\d+)\s*[-–—.]\s*")
ALTERNATIVE_PATTERN = re.compile(r"^([a-zA-Z])\)\s*")


def _get_text(p_element) -> str:
    runs = p_element.findall(".//w:r/w:t", NSMAP)
    return "".join(r.text or "" for r in runs)


def _has_image(p_element) -> bool:
    if p_element.findall(f".//{{{NS_W}}}drawing"):
        return True
    if p_element.findall(f".//{{{NS_MC}}}AlternateContent"):
        return True
    if p_element.findall(".//w:pict", NSMAP):
        return True
    for elem in p_element.iter():
        tag = elem.tag if isinstance(elem.tag, str) else ""
        if "drawing" in tag or "blip" in tag or "imagedata" in tag:
            return True
    return False


def _has_anchor(p_element) -> bool:
    return bool(p_element.findall(f".//{{{NS_WP}}}anchor"))


def _has_inline(p_element) -> bool:
    return bool(p_element.findall(f".//{{{NS_WP}}}inline"))


def _image_type(p_element) -> str:
    parts = []
    if _has_anchor(p_element):
        parts.append("ANCHOR(flotante)")
    if _has_inline(p_element):
        parts.append("INLINE")
    return ", ".join(parts) if parts else "desconocido"


def _find_questions_table(doc):
    if not doc.tables:
        raise ValueError("El documento no contiene tablas.")
    max_rows, best_idx = 0, 0
    for i, table in enumerate(doc.tables):
        if len(table.rows) > max_rows:
            max_rows = len(table.rows)
            best_idx = i
    return best_idx, doc.tables[best_idx]


def _find_cell_for_question(table, target_number: int, questions_per_row: int):
    """Recorre la tabla buscando la celda cuya primera línea sea la pregunta target."""
    for row_idx, row in enumerate(table.rows):
        for col_idx in range(questions_per_row):
            if col_idx >= len(row.cells):
                continue
            cell = row.cells[col_idx]
            tc = cell._tc
            p_elements = tc.findall("w:p", NSMAP)
            if not p_elements:
                continue
            first_text = _get_text(p_elements[0]).strip()
            m = QUESTION_PATTERN.match(first_text)
            if m and int(m.group(1)) == target_number:
                return row_idx, col_idx, cell, p_elements
    return None, None, None, None


def _classify_paragraph(p_elem, in_alternatives: bool, current_alt_letter) -> str:
    """Replica la lógica del parser para mostrar cómo clasificaría este párrafo."""
    text = _get_text(p_elem).strip()
    has_img = _has_image(p_elem)
    is_empty = not text and not has_img

    alt_match = ALTERNATIVE_PATTERN.match(text)
    q_match = QUESTION_PATTERN.match(text)

    if q_match:
        return f"ENUNCIADO (inicio pregunta {q_match.group(1)})"
    if alt_match:
        return f"ALTERNATIVA-INICIO ({alt_match.group(1)})"
    if is_empty and in_alternatives:
        return "TRAILING_EMPTY"
    if is_empty and not in_alternatives:
        return "ENUNCIADO (vacío/spacing)"
    if not in_alternatives:
        return "ENUNCIADO (continuación)"
    else:
        return f"ALT-CONTINUACIÓN (pertenece a '{current_alt_letter}')"


def inspect_question(file_path: str, question_number: int):
    print(f"\n{'=' * 70}")
    print(f"  Archivo : {file_path}")
    print(f"  Pregunta: {question_number}")
    print(f"{'=' * 70}\n")

    doc = Document(file_path)
    table_idx, table = _find_questions_table(doc)
    questions_per_row = len(table.columns) if table.columns else 2

    print(f"Tabla de preguntas: tabla[{table_idx}]  ({questions_per_row} col(s))\n")

    row_idx, col_idx, cell, p_elements = _find_cell_for_question(
        table, question_number, questions_per_row
    )

    if cell is None:
        print(f"❌  No se encontró la pregunta {question_number} en la tabla.")
        return

    print(f"Celda encontrada: fila {row_idx}, columna {col_idx}")
    print(f"Total párrafos en la celda: {len(p_elements)}\n")

    # Simular clasificación del parser
    in_alternatives = False
    current_alt_letter = None

    for i, p_elem in enumerate(p_elements):
        text = _get_text(p_elem)
        has_img = _has_image(p_elem)
        is_empty = not text.strip() and not has_img

        classification = _classify_paragraph(
            p_elem, in_alternatives, current_alt_letter
        )

        # Actualizar estado del simulador
        alt_match = ALTERNATIVE_PATTERN.match(text.strip())
        if alt_match:
            in_alternatives = True
            current_alt_letter = alt_match.group(1).lower()

        # Indicadores
        img_marker = f"  🖼  [{_image_type(p_elem)}]" if has_img else ""
        empty_marker = "  (párrafo vacío)" if is_empty else ""
        text_preview = repr(text[:80]) if text.strip() else ""

        print(f"  [{i:02d}] {classification}")
        if text_preview:
            print(f"        texto  : {text_preview}")
        if img_marker:
            print(f"        imagen : {img_marker.strip()}")
        if empty_marker:
            print(f"        {empty_marker.strip()}")

        # Mostrar XML resumido del párrafo (solo tags de nivel 1 dentro de <w:p>)
        children_tags = []
        for child in p_elem:
            short = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            children_tags.append(short)
        print(
            f"        xml    : <w:p> → {' | '.join(children_tags) if children_tags else '(vacío)'}"
        )
        print()

    # Dump XML completo de los párrafos que tienen imagen
    print(f"\n{'─' * 70}")
    print("  XML COMPLETO de los párrafos con imagen:")
    print(f"{'─' * 70}\n")
    for i, p_elem in enumerate(p_elements):
        if _has_image(p_elem):
            xml_str = etree.tostring(p_elem, pretty_print=True).decode()
            # Truncar si es muy largo
            lines = xml_str.splitlines()
            if len(lines) > 80:
                xml_str = (
                    "\n".join(lines[:80]) + f"\n  ... ({len(lines) - 80} líneas más)"
                )
            print(f"  --- Párrafo [{i:02d}] ---")
            print(xml_str)
            print()


def main():
    parser = argparse.ArgumentParser(
        description="Diagnóstico de estructura de párrafos para una pregunta específica."
    )
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--local", metavar="PATH", help="Ruta local al archivo .docx")
    group.add_argument(
        "--drive", metavar="DRIVE_PATH", help="Ruta en Google Drive al archivo .docx"
    )
    parser.add_argument(
        "--question",
        type=int,
        default=5,
        help="Número de pregunta a inspeccionar (default: 5)",
    )
    args = parser.parse_args()

    if args.local:
        file_path = args.local
    else:
        print("⬇️  Descargando desde Google Drive...")
        from src.drive_client import download_docx

        file_path = download_docx(args.drive)
        print(f"   Descargado en: {file_path}\n")

    inspect_question(file_path, args.question)


if __name__ == "__main__":
    main()
